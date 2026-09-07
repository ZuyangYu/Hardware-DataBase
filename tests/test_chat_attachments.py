"""Chat attachment domain tests (Phase 0/1/2 regression locks).

Covers: store idempotency, asset reuse, ACL, parse jobs, local document
parsing, FTS5 retrieval, the hardware query normalizer, and the turn waiting
state machine. The feature flag default-on behavior is also locked here.
"""

from __future__ import annotations

import io
import inspect
import os
import sqlite3
import tempfile
import unittest
import zipfile
from contextlib import closing

import src.settings
from src.agents.tools.attachment_tools import build_attachment_tools, make_attachment_list
from src.agents.tools.runtime import ToolRuntime
from src.attachments.coordinator import AttachmentTurnCoordinator
from src.attachments.index import AttachmentIndex, fts5_available
from src.attachments.jobs import AttachmentParseExecutor
from src.attachments.local_document_parser import LocalDocumentParser
from src.attachments.models import (
    AttachmentPart,
    AttachmentRef,
    AttachmentQuotaExceeded,
    PARSE_STATUS_FAILED,
    PARSE_STATUS_READY,
    PART_TYPE_IMAGE,
    PART_TYPE_OCR_TEXT,
    resolve_source_scope,
    resolve_auto_scope,
    scope_allows_attachments,
    scope_allows_kb,
)
from src.attachments.query_normalizer import (
    build_fts_match,
    extract_identifiers,
    fts_safe_term,
)
from src.attachments.retrieval import AttachmentRetrievalService
from src.attachments.service import AttachmentService
from src.attachments import storage
from src.attachments.store import AttachmentStore
from src.attachments.worker import AttachmentWorker
from src.core.conversation import ConversationService


def test_attachment_tools_can_be_registered_as_langchain_tools():
    """Every attachment callable must satisfy Deep Agent tool conversion."""
    from langchain_core.tools import StructuredTool

    runtime = ToolRuntime(
        kb_name="ADAS",
        ctx=None,
        source_scope="attachment_and_knowledge_base",
        attachment_refs=[
            AttachmentRef(
                attachment_id="att-1",
                asset_id="asset-1",
                session_id=1,
                filename="HSI.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                extension=".docx",
                size_bytes=1,
                sha256="hash",
            )
        ],
    )

    tools = build_attachment_tools(runtime)
    assert {tool.__name__ for tool in tools} == {
        "attachment_list",
        "attachment_search",
        "attachment_read",
        "attachment_table_query",
        "attachment_circuit_search",
        "attachment_visual_analyze",
    }
    for function in tools:
        assert inspect.getdoc(function)
        registered = StructuredTool.from_function(function)
        assert registered.name == function.__name__


def _minimal_xlsx_bytes() -> bytes:
    """Build a tiny OOXML workbook without adding an authoring dependency."""
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>'
        ),
        "xl/workbook.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Parts" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            'Target="worksheets/sheet1.xml"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>Reference</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>Value</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>R1</t></is></c>'
            '<c r="B2" t="inlineStr"><is><t>10k</t></is></c></row>'
            '</sheetData></worksheet>'
        ),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return output.getvalue()


def _empty_pdf_bytes() -> bytes:
    """Create a valid blank PDF page for OCR/degraded parser tests."""
    from reportlab.pdfgen import canvas

    output = io.BytesIO()
    document = canvas.Canvas(output)
    document.showPage()
    document.save()
    return output.getvalue()


class AttachmentDomainTestBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._patch(src.settings, "CHAT_ATTACHMENT_INDEX_DB_PATH", os.path.join(self.tmp.name, "att.db"))
        self._patch(src.settings, "CHAT_ATTACHMENT_STORAGE_DIR", os.path.join(self.tmp.name, "att-storage"))
        self._patch(src.settings, "CHAT_ATTACHMENTS_ENABLED", True)
        self._patch(src.settings, "AUTH_DB_PATH", os.path.join(self.tmp.name, "auth.db"))
        self.db_path = src.settings.AUTH_DB_PATH
        self._seed_user()
        self.store = AttachmentStore(db_path=src.settings.CHAT_ATTACHMENT_INDEX_DB_PATH)
        self.service = AttachmentService(store=self.store)
        self.conv = ConversationService(self.db_path)
        self.user_id = 1
        self.session = self.conv.create_session(self.user_id, "__general__")

    def _seed_user(self):
        import sqlite3
        from contextlib import closing

        with closing(sqlite3.connect(self.db_path)) as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT UNIQUE NOT NULL,
                    department_id INTEGER
                )
                """
            )
            conn.execute(
                "INSERT OR IGNORE INTO users (id, username, department_id) VALUES (1, 'user1', 1)"
            )
            conn.commit()

    def _patch(self, module, name, value):
        old = getattr(module, name)
        setattr(module, name, value)
        self.addCleanup(setattr, module, name, old)

    def _upload(self, filename: str, content: bytes, client_request_id: str | None = None):
        return self.service.upload(
            session_id=self.session.id,
            user_id=self.user_id,
            filename=filename,
            stream=io.BytesIO(content),
            client_request_id=client_request_id,
        )


class StoreTests(AttachmentDomainTestBase):
    def test_attachment_client_request_id_is_idempotent(self):
        first = self._upload("doc.txt", "hello".encode(), client_request_id="req-1")
        second = self._upload("doc.txt", "hello".encode(), client_request_id="req-1")
        self.assertEqual(first.attachment_id, second.attachment_id)

    def test_idempotent_retry_bypasses_session_file_quota(self):
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_FILES_PER_SESSION", 1)
        first = self._upload("doc.txt", b"hello", client_request_id="req-quota")

        retried = self._upload(
            "renamed.txt", b"different body", client_request_id="req-quota"
        )

        self.assertEqual(retried.attachment_id, first.attachment_id)
        self.assertEqual(self.store.count_active_attachments(
            session_id=self.session.id, user_id=self.user_id
        ), 1)

    def test_idempotent_retry_does_not_leave_an_unreferenced_asset(self):
        first = self._upload("doc.txt", b"hello", client_request_id="req-asset")

        retried = self._upload(
            "renamed.txt", b"different body", client_request_id="req-asset"
        )

        self.assertEqual(retried.attachment_id, first.attachment_id)
        self.assertIsNone(self.store.find_asset(
            tenant_id="default",
            user_id=self.user_id,
            session_id=self.session.id,
            sha256=storage.sha256_of_bytes(b"different body"),
        ))

    def test_session_and_tenant_byte_quotas_are_enforced(self):
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_SESSION_BYTES", 5)
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_TENANT_BYTES", 100)
        self._upload("first.txt", b"12345")
        with self.assertRaises(AttachmentQuotaExceeded):
            self._upload("second.txt", b"6")

    def test_tenant_byte_quota_applies_across_sessions(self):
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_SESSION_BYTES", 100)
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_TENANT_BYTES", 5)
        self._upload("first.txt", b"12345")
        other_session = self.conv.create_session(self.user_id, "__general__")
        with self.assertRaises(AttachmentQuotaExceeded):
            self.service.upload(
                session_id=other_session.id,
                user_id=self.user_id,
                filename="second.txt",
                stream=io.BytesIO(b"6"),
            )

    def test_same_content_in_session_shares_asset(self):
        first = self._upload("a.txt", "same-content".encode())
        second = self._upload("b.txt", "same-content".encode())
        self.assertNotEqual(first.attachment_id, second.attachment_id)
        self.assertEqual(first.asset_id, second.asset_id)

    def test_delete_last_reference_allows_same_hash_reupload_to_rebuild_asset(self):
        first = self._upload("board.txt", b"same-content")
        old_asset_id = first.asset_id
        self.assertTrue(
            self.service.delete_attachment(
                attachment_id=first.attachment_id,
                session_id=self.session.id,
                user_id=self.user_id,
            )
        )
        self.assertIsNone(self.store.get_asset(old_asset_id))

        second = self._upload("board.txt", b"same-content")
        self.assertNotEqual(second.asset_id, old_asset_id)
        self.assertEqual(self.store.get_asset(second.asset_id).parse_status, "queued")
        self.assertEqual(len(self.store.list_pending_jobs()), 1)

    def test_cross_session_dedup_is_not_enforced(self):
        other = self.conv.create_session(self.user_id, "__general__")
        first = self.service.upload(
            session_id=self.session.id, user_id=self.user_id,
            filename="a.txt", stream=io.BytesIO(b"same"), )
        second = self.service.upload(
            session_id=other.id, user_id=self.user_id,
            filename="a.txt", stream=io.BytesIO(b"same"), )
        self.assertNotEqual(first.asset_id, second.asset_id)

    def test_foreign_attachment_is_invisible_not_403(self):
        record = self._upload("doc.txt", "hello".encode())
        other = self.conv.create_session(self.user_id, "__general__")
        with self.assertRaises(Exception):
            self.service.get_attachment(
                attachment_id=record.attachment_id,
                session_id=other.id,
                user_id=self.user_id,
            )

    def test_delete_soft_deletes_and_keeps_history_safe(self):
        record = self._upload("doc.txt", "hello".encode())
        asset_id = record.asset_id
        self.assertTrue(
            self.service.delete_attachment(
                attachment_id=record.attachment_id,
                session_id=self.session.id,
                user_id=self.user_id,
            )
        )
        self.assertIsNone(
            self.store.get_attachment(attachment_id=record.attachment_id, include_deleted=False)
        )
        # Reference count dropped to zero -> asset cleanup removed parts.
        self.assertEqual(self.store.asset_reference_count(asset_id=asset_id), 0)

    def test_delete_last_reference_removes_trigram_index_rows(self):
        record = self._upload("board.txt", b"VDD_3V3 powers U2")
        parts = [
            AttachmentPart(
                part_id="part-trigram-delete",
                asset_id=record.asset_id,
                ordinal=0,
                part_type="text",
                text_content="VDD_3V3 powers U2",
            )
        ]
        self.store.replace_parts(record.asset_id, parts)
        index = AttachmentIndex(db_path=self.store.db_path)
        if not index.trigram_enabled:
            self.skipTest("SQLite trigram tokenizer is unavailable")
        index.replace_asset(record.asset_id, self.store.list_parts(record.asset_id))

        self.assertTrue(index.trigram_search("D_3V", asset_ids=[record.asset_id]))

        self.service.delete_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )

        self.assertEqual(index.trigram_search("D_3V", asset_ids=[record.asset_id]), [])

    def test_delete_asset_cleans_indexes_after_search_flags_are_disabled(self):
        record = self._upload("board.txt", b"VDD_3V3 powers U2")
        self.store.replace_parts(
            record.asset_id,
            [
                AttachmentPart(
                    part_id="part-flag-toggle-delete",
                    asset_id=record.asset_id,
                    ordinal=0,
                    part_type="text",
                    text_content="VDD_3V3 powers U2",
                )
            ],
        )
        index = AttachmentIndex(db_path=self.store.db_path)
        if not index.trigram_enabled:
            self.skipTest("SQLite trigram tokenizer is unavailable")
        index.replace_asset(record.asset_id, self.store.list_parts(record.asset_id))
        self.assertTrue(index.trigram_search("D_3V", asset_ids=[record.asset_id]))

        self._patch(src.settings, "CHAT_ATTACHMENT_FTS_ENABLED", False)
        self._patch(src.settings, "CHAT_ATTACHMENT_TRIGRAM_ENABLED", False)
        self.service.delete_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )

        self.assertEqual(index.trigram_search("D_3V", asset_ids=[record.asset_id]), [])

    def test_delete_one_of_two_attachments_keeps_shared_asset(self):
        first = self._upload("a.txt", "shared-bytes".encode())
        second = self._upload("b.txt", "shared-bytes".encode())
        self.service.delete_attachment(
            attachment_id=first.attachment_id, session_id=self.session.id, user_id=self.user_id
        )
        surviving = self.store.get_asset(second.asset_id)
        self.assertIsNotNone(surviving)

    def test_session_cleanup_outbox_is_idempotent(self):
        created = self.store.enqueue_session_cleanup(
            session_id=self.session.id, user_id=self.user_id
        )
        self.assertTrue(created)
        self.assertFalse(self.store.enqueue_session_cleanup(
            session_id=self.session.id, user_id=self.user_id
        ))
        pending = self.store.list_pending_cleanups()
        self.assertEqual(len(pending), 1)

    def test_session_cleanup_coordinator_reclaims_result_exports(self):
        from src.core.session_cleanup import SessionCleanupCoordinator
        from src.result_exports.models import ResultEnvelope
        from src.result_exports.store import ResultExportStore
        from src.result_exports.worker import ResultExportWorker

        result_store = ResultExportStore(
            self.db_path, storage_dir=os.path.join(self.tmp.name, "exports")
        )
        snapshot = result_store.create_snapshot(
            owner_user_id=self.user_id,
            tenant_id="default",
            session_id=self.session.id,
            turn_id="turn-cleanup",
            envelope=ResultEnvelope(answer="answer"),
        )
        job = result_store.create_export_job(
            owner_user_id=self.user_id,
            tenant_id="default",
            session_id=self.session.id,
            snapshot_id=snapshot.snapshot_id,
            format="md",
            client_request_id="cleanup-result",
        )
        self.assertTrue(ResultExportWorker(store=result_store, worker_id="export-w").run_once())
        artifact = result_store.get_export_job(self.user_id, job.export_job_id)
        self.assertIsNotNone(artifact)
        artifact_row = result_store.get_artifact(self.user_id, artifact.artifact_id)
        self.assertIsNotNone(artifact_row)
        artifact_path = result_store.storage_dir / artifact_row.storage_ref
        self.assertTrue(artifact_path.exists())
        self.store.enqueue_session_cleanup(session_id=self.session.id, user_id=self.user_id)

        coordinator = SessionCleanupCoordinator(
            store=self.store, result_store=result_store, worker_id="cleanup-w"
        )
        self.assertTrue(coordinator.run_once())

        self.assertFalse(artifact_path.exists())
        self.assertIsNone(result_store.get_snapshot(self.user_id, snapshot.snapshot_id))
        self.assertIsNone(result_store.get_export_job(self.user_id, job.export_job_id))

    def test_session_cleanup_coordinator_uses_default_result_store(self):
        from src.core.session_cleanup import SessionCleanupCoordinator
        from src.result_exports.models import ResultEnvelope
        from src.result_exports.store import ResultExportStore
        from src.result_exports.worker import ResultExportWorker

        export_dir = os.path.join(self.tmp.name, "default-exports")
        self._patch(src.settings, "RESULT_EXPORT_STORAGE_DIR", export_dir)
        result_store = ResultExportStore(self.db_path)
        snapshot = result_store.create_snapshot(
            owner_user_id=self.user_id,
            tenant_id="default",
            session_id=self.session.id,
            turn_id="turn-default-cleanup",
            envelope=ResultEnvelope(answer="answer"),
        )
        job = result_store.create_export_job(
            owner_user_id=self.user_id,
            tenant_id="default",
            session_id=self.session.id,
            snapshot_id=snapshot.snapshot_id,
            format="md",
            client_request_id="default-cleanup-result",
        )
        ResultExportWorker(store=result_store, worker_id="export-w").run_once()
        completed = result_store.get_export_job(self.user_id, job.export_job_id)
        artifact = result_store.get_artifact(self.user_id, completed.artifact_id)
        artifact_path = result_store.storage_dir / artifact.storage_ref
        self.store.enqueue_session_cleanup(session_id=self.session.id, user_id=self.user_id)

        self.assertTrue(SessionCleanupCoordinator(store=self.store, worker_id="cleanup-w").run_once())

        self.assertFalse(artifact_path.exists())
        self.assertIsNone(result_store.get_snapshot(self.user_id, snapshot.snapshot_id))

    def test_session_cleanup_coordinator_reclaims_document_resources(self):
        from unittest.mock import Mock, patch

        from src.core.session_cleanup import SessionCleanupCoordinator

        document_jobs = Mock()
        document_jobs.cancel_session.return_value = ["wo-chat-1"]
        document_store = Mock()
        result_store = Mock()
        self.store.enqueue_session_cleanup(session_id=self.session.id, user_id=self.user_id)

        coordinator = SessionCleanupCoordinator(
            store=self.store,
            result_store=result_store,
            document_job_store=document_jobs,
            document_store=document_store,
            worker_id="cleanup-w",
        )
        with patch("src.agents.runner.forget_thread", return_value=True) as forget_thread:
            self.assertTrue(coordinator.run_once())

        document_jobs.cancel_session.assert_called_once_with(
            session_id=self.session.id, reason="chat session deleted"
        )
        document_store.cleanup_work_orders.assert_called_once_with(["wo-chat-1"])
        document_jobs.cleanup_session.assert_called_once_with(session_id=self.session.id)
        result_store.cleanup_session.assert_called_once_with(session_id=self.session.id)
        forget_thread.assert_called_once_with(str(self.session.id))


class JobTests(AttachmentDomainTestBase):
    def test_claiming_parse_job_marks_asset_running(self):
        record = self._upload("notes.txt", b"content")
        job = self.store.list_pending_jobs()[0]

        claimed = self.store.claim_job(job.job_id, "w1")

        self.assertIsNotNone(claimed)
        self.assertEqual(self.store.get_asset(record.asset_id).parse_status, "running")

    def test_wrong_worker_cannot_complete_job_or_settle_asset(self):
        record = self._upload("notes.txt", b"content")
        job = self.store.list_pending_jobs()[0]
        self.assertIsNotNone(self.store.claim_job(job.job_id, "owner-worker"))

        self.store.complete_job(
            job.job_id,
            "wrong-worker",
            {"parse_status": PARSE_STATUS_READY},
        )

        self.assertEqual(self.store.get_job(job.job_id).status, "running")
        self.assertEqual(self.store.get_asset(record.asset_id).parse_status, "running")

    def test_failed_parse_job_updates_asset_for_retry_and_dead_letter(self):
        record = self._upload("notes.txt", b"content")
        job = self.store.list_pending_jobs()[0]

        for attempt in range(job.max_attempts):
            worker_id = f"failure-worker-{attempt}"
            self.assertIsNotNone(self.store.claim_job(job.job_id, worker_id))
            self.store.fail_job(job.job_id, worker_id, "worker_error", "temporary failure")
            expected = "failed" if attempt == job.max_attempts - 1 else "queued"
            self.assertEqual(self.store.get_asset(record.asset_id).parse_status, expected)

    def test_parse_job_completes_and_resolves_state(self):
        record = self._upload("notes.txt", "STM32H743 powers the board".encode())
        worker = AttachmentWorker(store=self.store, worker_id="w1")
        self.assertTrue(worker.run_once())
        asset = self.store.get_asset(record.asset_id)
        self.assertEqual(asset.parse_status, PARSE_STATUS_READY)
        self.assertIn("STM32H743", "\n".join(p.text_content for p in self.store.list_parts(record.asset_id)))
        # No pending work the second time.
        self.assertFalse(worker.run_once())

    def test_rejected_extension_fails_closed(self):
        with self.assertRaises(Exception):
            self._upload("legacy.xls", b"binary")

    def test_parse_failure_is_recorded(self):
        record = self._upload("broken.pdf", b"%PDF-1.4 not really a pdf")
        worker = AttachmentWorker(store=self.store, worker_id="w1")
        worker.run_once()
        asset = self.store.get_asset(record.asset_id)
        self.assertIn(asset.parse_status, {PARSE_STATUS_FAILED, PARSE_STATUS_READY})

    def test_retry_parse_resets_failed_asset_to_queued(self):
        record = self._upload("notes.txt", b"content")
        job = self.store.list_pending_jobs()[0]
        for attempt in range(job.max_attempts):
            claimed = self.store.claim_job(job.job_id, f"failed-worker-{attempt}")
            self.assertIsNotNone(claimed)
            self.store.fail_job(job.job_id, f"failed-worker-{attempt}", "parse_failed", "boom")

        self.store.update_asset_status(
            record.asset_id,
            parse_status=PARSE_STATUS_FAILED,
            error_code="parse_failed",
            error_message="boom",
        )
        retried = self.service.retry_parse(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )

        self.assertEqual(retried.parse_status, "queued")
        self.assertEqual(self.store.get_asset(record.asset_id).parse_status, "queued")
        self.assertEqual(self.store.get_asset(record.asset_id).error_code, "")
        self.assertEqual(len(self.store.list_pending_jobs()), 1)

    def test_unresolvable_storage_key_marks_asset_failed_and_fails_waiting_turn(self):
        record = self._upload("notes.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id,
            self.session.id,
            "分析附件",
            attachment_ids=[record.attachment_id],
        )
        with self.store._connect() as conn:
            conn.execute(
                "UPDATE chat_attachment_assets SET storage_key = ? WHERE asset_id = ?",
                ("../outside/source.txt", record.asset_id),
            )

        worker = AttachmentWorker(store=self.store, worker_id="invalid-source-worker")
        self.assertTrue(worker.run_once())

        asset = self.store.get_asset(record.asset_id)
        self.assertEqual(asset.parse_status, PARSE_STATUS_FAILED)
        self.assertEqual(asset.error_code, "source_unavailable")
        self.assertEqual(self.conv.get_turn_unscoped(turn.id).status, "failed")


class ParsingAndRetrievalTests(AttachmentDomainTestBase):
    def _parsed_txt_asset(self, text: str):
        record = self._upload("board.md", text.encode())
        executor = AttachmentParseExecutor(store=self.store)
        result = executor.process_asset(record.asset_id)
        self.assertTrue(result["ok"], result)
        return record

    def test_query_normalizer_hardware_identifiers(self):
        identifiers = extract_identifiers("对比 STM32H743ZI 和 TPS62130 的 VDD_3V3 供电，R1 的阻值")
        lowered = [item.lower() for item in identifiers]
        self.assertIn("stm32h743zi", lowered)
        self.assertIn("tps62130", lowered)
        self.assertIn("vdd_3v3", lowered)
        self.assertIn("r1", lowered)

    def test_query_normalizer_fts_never_raises(self):
        for query in ["+12V", "-5V rail", "A/B 测试", "R1/R2", "STM32H7*", "CAN_H/CAN_L", "(select", "OR 1=1"]:
            match_expr = build_fts_match(query)
            safe = fts_safe_term(query)
            self.assertIsInstance(match_expr, str)
            self.assertIsInstance(safe, str)

    def test_lexical_retrieval_finds_identifier_and_cjk(self):
        text = (
            "# 电源设计\n"
            "主电源采用 TPS62130，输出 VDD_3V3 网络，最大电流 3A。\n"
            "以太网 PHY 使用 ETH_TXP 与 ETH_TXN 差分对连接到 RJ45。\n"
            "连续中文段落用于验证分词缺失时的回退检索能力。\n"
        )
        record = self._parsed_txt_asset(text)
        retrieval = AttachmentRetrievalService()
        result = retrieval.search(
            "TPS62130 输出电压",
            asset_ids=[record.asset_id],
            limit=5,
        )
        self.assertTrue(result.chunks)
        joined = "\n".join(chunk.text_content for chunk in result.chunks)
        self.assertIn("TPS62130", joined)

    def test_trigram_lookup_returns_substring_hits(self):
        record = self._parsed_txt_asset("Power rail VDD_3V3 feeds the PHY.")
        index = AttachmentIndex(db_path=self.store.db_path)
        if not index.trigram_enabled:
            self.skipTest("SQLite trigram tokenizer is unavailable")
        hits = index.trigram_search("D_3V", asset_ids=[record.asset_id], limit=5)
        self.assertTrue(hits)
        self.assertEqual(hits[0].asset_id, record.asset_id)

    def test_fts_disabled_falls_back_to_canonical_parts(self):
        record = self._parsed_txt_asset("R1 is a 10k pull-up on VDD_3V3.")
        self._patch(src.settings, "CHAT_ATTACHMENT_FTS_ENABLED", False)
        self._patch(src.settings, "CHAT_ATTACHMENT_TRIGRAM_ENABLED", False)
        index = AttachmentIndex(db_path=self.store.db_path)
        result = AttachmentRetrievalService(index=index).search(
            "R1 pull-up", asset_ids=[record.asset_id], limit=5
        )
        self.assertTrue(result.chunks)
        self.assertIn("scan", result.matched_backends)
        self.assertIn("fts5_unavailable", result.degraded_reasons)

    def test_identifier_variants_match_separator_changes(self):
        record = self._parsed_txt_asset("The VDD_3V3 rail is generated by the regulator.")
        result = AttachmentRetrievalService().search(
            "vdd-3v3", asset_ids=[record.asset_id], limit=5
        )
        self.assertTrue(result.chunks)
        self.assertIn("VDD_3V3", "\n".join(chunk.text_content for chunk in result.chunks))

    def test_exact_refdes_lookup(self):
        text = "R1 connects VDD_3V3 to the LED. R2 is a pull-up.\n"
        record = self._parsed_txt_asset(text)
        retrieval = AttachmentRetrievalService()
        result = retrieval.search("R1 阻值", asset_ids=[record.asset_id], limit=5)
        self.assertTrue(result.chunks)
        self.assertIn("R1", "\n".join(chunk.text_content for chunk in result.chunks))

    def test_context_ratio_caps_default_attachment_budget(self):
        record = self._parsed_txt_asset("R1 " + ("long attachment evidence " * 60))
        self._patch(src.settings, "CHAT_ATTACHMENT_CONTEXT_MAX_TOKENS", 1000)
        self._patch(src.settings, "CHAT_ATTACHMENT_CONTEXT_RATIO", 0.1)
        self._patch(src.settings, "AGENT_MODEL_MAX_INPUT_TOKENS", 1000)

        result = AttachmentRetrievalService().search(
            "R1", asset_ids=[record.asset_id], limit=5
        )

        self.assertTrue(result.truncated_by_budget)
        self.assertEqual(result.chunks, [])

    def test_fts5_probe_reports_availability(self):
        self.assertTrue(fts5_available(self.store.db_path))

    def test_pdf_scan_degrades_without_fabricating_text(self):
        # A minimal valid PDF with one empty page -> degraded, no text parts.
        pdf = _empty_pdf_bytes()
        record = self._upload("empty.pdf", pdf)
        executor = AttachmentParseExecutor(store=self.store)
        result = executor.process_asset(record.asset_id)
        if result["ok"]:
            self.assertEqual(result["parse_status"], "degraded")
            self.assertTrue(result["manifest"].get("ocr_candidate"))
            # Scanned pages must not fabricate text evidence.
            self.assertEqual(result["part_count"], 0)

    def test_docx_parser_exposes_links_and_embedded_image_metadata(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from PIL import Image

        image = io.BytesIO()
        Image.new("RGB", (1, 1), "white").save(image, format="PNG")
        document = Document()
        paragraph = document.add_paragraph("Hardware reference")
        document.add_picture(io.BytesIO(image.getvalue()))
        relationship_id = document.part.relate_to(
            "https://example.test/reference", RT.HYPERLINK, is_external=True
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "external reference"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        source = io.BytesIO()
        document.save(source)

        record = self._upload("reference.docx", source.getvalue())
        result = AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)

        self.assertTrue(result["ok"], result)
        self.assertEqual(result["manifest"]["link_count"], 1)
        self.assertEqual(result["manifest"]["embedded_image_count"], 1)
        parts = self.store.list_parts(record.asset_id)
        self.assertTrue(any(part.part_type == PART_TYPE_IMAGE for part in parts))

    def test_pdf_ocr_disabled_does_not_call_engine_or_fabricate_text(self):
        pdf = _empty_pdf_bytes()
        record = self._upload("scanned.pdf", pdf)

        class RecordingOcrEngine:
            provider = "test_cpu"

            def __init__(self):
                self.calls = []

            def recognize_page(self, source_path, page_number):
                self.calls.append((source_path, page_number))
                return "should not be used"

        engine = RecordingOcrEngine()
        self._patch(src.settings, "CHAT_ATTACHMENT_OCR_ENABLED", False)
        outcome = LocalDocumentParser(ocr_engine=engine).parse(
            self.store.get_asset(record.asset_id),
            storage.resolve_storage_key(self.store.get_asset(record.asset_id).storage_key),
        )

        self.assertTrue(outcome.ok)
        self.assertEqual(outcome.parse_status, "degraded")
        self.assertEqual(engine.calls, [])
        self.assertFalse(any(part.part_type == PART_TYPE_OCR_TEXT for part in outcome.parts))

    def test_pdf_ocr_success_appends_page_located_ocr_text_and_clears_degraded(self):
        pdf = _empty_pdf_bytes()
        record = self._upload("scanned.pdf", pdf)

        class RecordingOcrEngine:
            provider = "test_cpu"

            def __init__(self):
                self.calls = []

            def recognize_page(self, source_path, page_number):
                self.calls.append((source_path, page_number))
                return "TPS62130 VDD_3V3"

        engine = RecordingOcrEngine()
        self._patch(src.settings, "CHAT_ATTACHMENT_OCR_ENABLED", True)
        outcome = LocalDocumentParser(ocr_engine=engine).parse(
            self.store.get_asset(record.asset_id),
            storage.resolve_storage_key(self.store.get_asset(record.asset_id).storage_key),
        )

        self.assertTrue(outcome.ok)
        self.assertEqual(outcome.parse_status, PARSE_STATUS_READY)
        self.assertEqual(len(engine.calls), 1)
        ocr_parts = [part for part in outcome.parts if part.part_type == PART_TYPE_OCR_TEXT]
        self.assertEqual(len(ocr_parts), 1)
        self.assertEqual(ocr_parts[0].text_content, "TPS62130 VDD_3V3")
        self.assertEqual(ocr_parts[0].locator["page"], 1)
        self.assertEqual(ocr_parts[0].metadata["provider"], "test_cpu")
        self.assertEqual(outcome.manifest["ocr_indexed_pages"], [1])

    def test_pdf_ocr_failure_keeps_degraded_without_fabricating_text(self):
        pdf = _empty_pdf_bytes()
        record = self._upload("scanned.pdf", pdf)

        class FailingOcrEngine:
            provider = "test_cpu"

            def recognize_page(self, source_path, page_number):
                raise RuntimeError("tesseract unavailable")

        self._patch(src.settings, "CHAT_ATTACHMENT_OCR_ENABLED", True)
        outcome = LocalDocumentParser(ocr_engine=FailingOcrEngine()).parse(
            self.store.get_asset(record.asset_id),
            storage.resolve_storage_key(self.store.get_asset(record.asset_id).storage_key),
        )

        self.assertTrue(outcome.ok)
        self.assertEqual(outcome.parse_status, "degraded")
        self.assertFalse(any(part.part_type == PART_TYPE_OCR_TEXT for part in outcome.parts))
        self.assertEqual(outcome.manifest["ocr_failed_pages"], [1])
        self.assertTrue(any("OCR" in reason for reason in outcome.degraded_reasons))

    def test_upload_records_mime_and_rejects_signature_mismatch(self):
        text = self._upload("notes.txt", b"plain text")
        self.assertEqual(text.media_type, "text/plain")
        with self.assertRaises(Exception):
            self.service.upload(
                session_id=self.session.id,
                user_id=self.user_id,
                filename="fake.pdf",
                stream=io.BytesIO(b"this is not a PDF"),
            )
        with self.assertRaises(Exception):
            self.service.upload(
                session_id=self.session.id,
                user_id=self.user_id,
                filename="notes.txt",
                stream=io.BytesIO(b"plain text"),
                declared_media_type="application/pdf",
            )

    def test_structured_attachment_paths_are_root_relative_and_readable(self):
        record = self._upload("parts.xlsx", _minimal_xlsx_bytes())
        result = AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)
        self.assertTrue(result["ok"], result)
        index_key = result["manifest"]["index_db_path"]
        self.assertFalse(os.path.isabs(index_key))
        resolved_index = storage.resolve_storage_key(index_key)
        self.assertTrue(os.path.isfile(resolved_index))

        from src.agents.tools.attachment_tools import make_attachment_table_query
        from src.agents.tools.runtime import ToolRuntime

        ref = self.service.build_refs([self.store.get_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )])[0]
        runtime = ToolRuntime(
            kb_name="",
            ctx=None,
            chat_session_id=str(self.session.id),
            attachment_refs=[ref],
            source_scope="attachment_only",
            attachment_service=self.service,
            attachment_user_id=self.user_id,
            attachment_session_id=self.session.id,
        )
        output = make_attachment_table_query(runtime)(query="")
        self.assertNotIn("没有可查询的 Excel", output)
        self.assertIn("Reference", output)

    def test_attachment_table_query_formats_sql_rows_and_keeps_record_scope(self):
        record = self._upload("parts.xlsx", _minimal_xlsx_bytes())
        result = AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)
        self.assertTrue(result["ok"], result)
        index_db = storage.resolve_storage_key(result["manifest"]["index_db_path"])
        with closing(sqlite3.connect(index_db)) as conn:
            table_name = conn.execute(
                "SELECT table_name FROM sql_table_registry WHERE record_id = ?",
                (result["manifest"]["record_id"],),
            ).fetchone()[0]

        from src.agents.tools.attachment_tools import make_attachment_table_query
        from src.agents.tools.runtime import ToolRuntime

        ref = self.service.build_refs([self.store.get_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )])[0]
        runtime = ToolRuntime(
            kb_name="",
            ctx=None,
            chat_session_id=str(self.session.id),
            attachment_refs=[ref],
            source_scope="attachment_only",
            attachment_service=self.service,
            attachment_user_id=self.user_id,
            attachment_session_id=self.session.id,
        )
        output = make_attachment_table_query(runtime)(sql=f'SELECT * FROM "{table_name}"')
        self.assertIn("col_1=R1", output)
        self.assertIn("R1", output)

    def test_xlsm_attachment_uses_read_only_structured_parser(self):
        record = self._upload("parts.xlsm", _minimal_xlsx_bytes())
        result = AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["parse_status"], PARSE_STATUS_READY)
        self.assertEqual(result["manifest"]["format"], "xlsx")

    def test_xlsx_row_limit_is_enforced_and_recorded_as_parse_failure(self):
        record = self._upload("parts.xlsx", _minimal_xlsx_bytes())
        self._patch(src.settings, "CHAT_ATTACHMENT_MAX_ROWS", 1)
        result = AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "parse_failed")
        self.assertEqual(self.store.get_asset(record.asset_id).parse_status, PARSE_STATUS_FAILED)

    def test_circuit_manifest_root_is_root_relative(self):
        from src.attachments.jobs import CircuitProcessor

        record = self._upload("board.edf", b"(edif board)")
        asset = self.store.get_asset(record.asset_id)
        outcome = CircuitProcessor().parse(asset, storage.resolve_storage_key(asset.storage_key))
        circuit_key = outcome.manifest["circuit_root"]
        self.assertFalse(os.path.isabs(circuit_key))
        self.assertEqual(
            os.path.abspath(storage.resolve_storage_key(circuit_key)),
            os.path.abspath(storage.asset_circuit_root(self.session.id)),
        )


class ScopeResolutionTests(AttachmentDomainTestBase):
    def test_auto_scope_matrix(self):
        self.assertEqual(resolve_auto_scope(has_attachments=False, is_general_chat=True), "knowledge_base_only")
        self.assertEqual(resolve_auto_scope(has_attachments=True, is_general_chat=True), "attachment_only")
        self.assertEqual(resolve_auto_scope(has_attachments=False, is_general_chat=False), "knowledge_base_only")
        self.assertEqual(resolve_auto_scope(has_attachments=True, is_general_chat=False), "attachment_and_knowledge_base")

    def test_scope_helpers(self):
        self.assertTrue(scope_allows_attachments("attachment_only"))
        self.assertFalse(scope_allows_kb("attachment_only"))
        self.assertTrue(scope_allows_kb("attachment_and_knowledge_base"))

    def test_explicit_scope_can_narrow_but_never_widens(self):
        self.assertEqual(
            resolve_source_scope(
                requested_scope="knowledge_base_only",
                has_attachments=True,
                is_general_chat=False,
            ),
            "knowledge_base_only",
        )
        self.assertEqual(
            resolve_source_scope(
                requested_scope="attachment_and_knowledge_base",
                has_attachments=True,
                is_general_chat=True,
            ),
            "attachment_only",
        )
        self.assertEqual(
            resolve_source_scope(
                requested_scope="attachment_and_knowledge_base",
                has_attachments=False,
                is_general_chat=False,
            ),
            "knowledge_base_only",
        )

    def test_attachment_tools_are_not_mounted_for_knowledge_base_only(self):
        record = self._upload("notes.txt", b"content")
        ref = self.service.build_refs([record])[0]
        runtime = ToolRuntime(
            kb_name="shared",
            ctx=None,
            attachment_refs=[ref],
            source_scope="knowledge_base_only",
        )
        self.assertEqual(build_attachment_tools(runtime), [])

    def test_feature_flag_defaults_on(self):
        self.assertEqual(src.settings.DEFAULT_VALUES.get("CHAT_ATTACHMENTS_ENABLED"), "true")

    def test_xlsx_limits_are_opt_in_and_do_not_change_default_kb_parser(self):
        from src.pipelines.spreadsheet.xlsx_parser import parse_xlsx

        workbook = parse_xlsx(io.BytesIO(_minimal_xlsx_bytes()))
        self.assertEqual(workbook.sheets[0].rows[0], ["Reference", "Value"])

        with self.assertRaises(ValueError):
            parse_xlsx(
                io.BytesIO(_minimal_xlsx_bytes()),
                limits={"max_rows": 1},
            )


class WaitingTurnTests(AttachmentDomainTestBase):
    def test_knowledge_base_only_drops_requested_attachment_capability(self):
        kb_session = self.conv.create_session(self.user_id, "shared")
        record = self.service.upload(
            session_id=kb_session.id,
            user_id=self.user_id,
            filename="notes.txt",
            stream=io.BytesIO(b"content"),
        )
        turn = self.conv.create_turn(
            self.user_id,
            kb_session.id,
            "只查询知识库",
            attachment_ids=[record.attachment_id],
            source_scope="knowledge_base_only",
        )
        self.assertEqual(turn.status, "pending")
        self.assertEqual(turn.source_scope, "knowledge_base_only")
        self.assertEqual(turn.required_attachment_ids, [])
        self.assertEqual(self.conv.list_turn_attachments(turn.id), [])
        self.assertEqual(turn.attachments, [])

    def test_turn_waits_then_resolves_to_pending(self):
        # Upload without running the worker -> parse still queued.
        record = self._upload("notes.txt", "hello world".encode())
        self.assertEqual(self.store.get_asset(record.asset_id).parse_status, "queued")
        turn = self.conv.create_turn(
            self.user_id,
            self.session.id,
            "总结这个附件",
            attachment_ids=[record.attachment_id],
            source_scope="auto",
        )
        self.assertEqual(turn.status, "waiting_for_attachments")
        snapshots = self.conv.list_turn_attachments(turn.id)
        self.assertEqual(len(snapshots), 1)
        self.assertEqual(snapshots[0]["attachment_id"], record.attachment_id)

        # Ready asset -> coordinator flips the turn to pending. Scope stored
        # at turn creation is persisted as-is; the API layer resolves "auto"
        # before calling create_turn, and the runner expands it defensively.
        executor = AttachmentParseExecutor(store=self.store)
        executor.process_asset(record.asset_id)
        coordinator = AttachmentTurnCoordinator(store=self.store)
        self.assertEqual(coordinator.notify_asset_settled(record.asset_id), 1)
        resolved = self.conv.get_turn_unscoped(turn.id)
        self.assertEqual(resolved.status, "pending")

    def test_failed_attachment_fails_the_waiting_turn(self):
        record = self._upload("doc.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.assertEqual(turn.status, "waiting_for_attachments")
        # Force the asset into a failed state as a broken parse would.
        self.store.update_asset_status(
            record.asset_id, parse_status=PARSE_STATUS_FAILED,
            error_code="parse_failed", error_message="boom",
        )
        coordinator = AttachmentTurnCoordinator(store=self.store)
        coordinator.notify_asset_settled(record.asset_id)
        failed = self.conv.get_turn_unscoped(turn.id)
        self.assertEqual(failed.status, "failed")

    def test_blocked_attachments_reject_turn_creation(self):
        record = self._upload("doc.txt", b"content")
        foreign_session = self.conv.create_session(self.user_id, "__general__")
        with self.assertRaises(ValueError):
            self.conv.create_turn(
                self.user_id, foreign_session.id, "试试",
                attachment_ids=[record.attachment_id],
            )

    def test_cancel_waiting_turn(self):
        record = self._upload("doc.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        cancelled = self.conv.request_turn_cancel(self.user_id, turn.id)
        self.assertEqual(cancelled.status, "cancelled")
        # Coordinator must not resurrect a cancelled turn.
        executor = AttachmentParseExecutor(store=self.store)
        executor.process_asset(record.asset_id)
        coordinator = AttachmentTurnCoordinator(store=self.store)
        coordinator.notify_asset_settled(record.asset_id)
        self.assertEqual(self.conv.get_turn_unscoped(turn.id).status, "cancelled")

    def test_deleting_required_attachment_fails_waiting_turn(self):
        record = self._upload("doc.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.assertEqual(turn.status, "waiting_for_attachments")

        self.service.delete_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )

        resolved = self.conv.get_turn_unscoped(turn.id)
        self.assertEqual(resolved.status, "failed")
        self.assertIn("附件", resolved.error_message)

    def test_historical_attachment_snapshot_marks_deleted_attachment(self):
        record = self._upload("design-review.pdf", b"%PDF-1.4\n")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )

        self.service.delete_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )

        snapshot = self.conv.list_turn_attachments(turn.id)[0]
        self.assertTrue(snapshot["deleted"])
        self.assertEqual(snapshot["filename_snapshot"], "design-review.pdf")

    def test_attachment_snapshot_projects_current_parse_status_while_active(self):
        record = self._upload("design-review.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.assertEqual(
            self.conv.list_turn_attachments(turn.id)[0]["parse_status_snapshot"],
            "queued",
        )

        self.store.update_asset_status(record.asset_id, parse_status=PARSE_STATUS_READY)

        snapshot = self.conv.list_turn_attachments(turn.id)[0]
        self.assertEqual(snapshot["parse_status_snapshot"], PARSE_STATUS_READY)
        self.assertFalse(snapshot["deleted"])

    def test_ready_attachment_creates_pending_turn_immediately(self):
        record = self._upload("doc.txt", b"content")
        executor = AttachmentParseExecutor(store=self.store)
        executor.process_asset(record.asset_id)
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.assertEqual(turn.status, "pending")

    def test_turn_snapshot_keeps_attachment_parser_version(self):
        record = self._upload("doc.txt", b"content")
        AttachmentParseExecutor(store=self.store).process_asset(record.asset_id)
        asset = self.store.get_asset(record.asset_id)

        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )

        snapshot = self.conv.list_turn_attachments(turn.id)[0]
        self.assertTrue(asset.parser_version)
        self.assertEqual(snapshot["parser_version_snapshot"], asset.parser_version)

    def test_reconciliation_resolves_waiting_turn_after_missed_settlement_event(self):
        record = self._upload("doc.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.assertEqual(turn.status, "waiting_for_attachments")

        # Simulate a worker restart after the asset was settled but before the
        # coordinator callback was delivered.
        self.store.update_asset_status(record.asset_id, parse_status=PARSE_STATUS_READY)

        coordinator = AttachmentTurnCoordinator(store=self.store)
        self.assertEqual(coordinator.reconcile_waiting_turns(), 1)
        self.assertEqual(self.conv.get_turn_unscoped(turn.id).status, "pending")

    def test_attachment_worker_runs_waiting_turn_reconciliation(self):
        record = self._upload("doc.txt", b"content")
        turn = self.conv.create_turn(
            self.user_id, self.session.id, "分析", attachment_ids=[record.attachment_id]
        )
        self.store.update_asset_status(record.asset_id, parse_status=PARSE_STATUS_READY)
        job = self.store.list_pending_jobs()[0]
        claimed = self.store.claim_job(job.job_id, "previous-worker")
        self.assertIsNotNone(claimed)
        self.store.complete_job(job.job_id, "previous-worker", {"parse_status": PARSE_STATUS_READY})

        worker = AttachmentWorker(store=self.store, worker_id="recovery-worker")
        self.assertTrue(worker.run_once())
        self.assertEqual(self.conv.get_turn_unscoped(turn.id).status, "pending")


class AttachmentToolAclTests(AttachmentDomainTestBase):
    def test_attachment_list_rechecks_user_and_session_on_every_call(self):
        record = self._upload("notes.txt", b"content")
        ref = self.service.build_refs([record])[0]
        runtime = ToolRuntime(
            kb_name="",
            ctx=None,
            chat_session_id=str(self.session.id),
            attachment_refs=[ref],
            source_scope="attachment_only",
            attachment_service=self.service,
            attachment_user_id=self.user_id,
            attachment_session_id=self.session.id,
        )
        tool = make_attachment_list(runtime)
        self.assertIn("notes.txt", tool())

        self.service.delete_attachment(
            attachment_id=record.attachment_id,
            session_id=self.session.id,
            user_id=self.user_id,
        )
        self.assertEqual(tool(), "本轮会话没有挂载任何附件。")

        foreign_runtime = ToolRuntime(
            kb_name="",
            ctx=None,
            chat_session_id=str(self.session.id),
            attachment_refs=[ref],
            source_scope="attachment_only",
            attachment_service=self.service,
            attachment_user_id=999,
            attachment_session_id=self.session.id,
        )
        self.assertEqual(make_attachment_list(foreign_runtime)(), "本轮会话没有挂载任何附件。")


class RegistryRegressionTests(unittest.TestCase):
    """Lock current KB invariants the attachment feature must not break."""

    def test_pipeline_registry_extension_conflict_still_raises(self):
        from src.pipelines.registry import PipelineRegistry, PipelineSpec

        spec_a = PipelineSpec(
            key="a", label="A", processor_kind="doc", content_kind="text",
            supported_extensions=frozenset({".pdf"}),
        )
        spec_b = PipelineSpec(
            key="b", label="B", processor_kind="doc", content_kind="text",
            supported_extensions=frozenset({".pdf"}),
        )
        registry = PipelineRegistry()
        registry.register(spec_a)
        with self.assertRaises(ValueError):
            registry.register(spec_b)

    def test_worker_queue_only_consumes_pending(self):
        import sqlite3
        from contextlib import closing

        db_path = os.path.join(tempfile.mkdtemp(), "auth.db")
        with closing(sqlite3.connect(db_path)) as conn:
            conn.execute(
                "CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY, username TEXT, department_id INTEGER, is_active INTEGER NOT NULL DEFAULT 1)"
            )
            conn.execute(
                "INSERT OR IGNORE INTO users (id, username, department_id) VALUES (1, 'user1', 1)"
            )
            conn.commit()
        conv = ConversationService(db_path)
        user_id = 1
        session = conv.create_session(user_id, "__general__")
        turn = conv.create_turn(user_id, session.id, "hello")
        pending = conv.list_pending_turn_work(limit=8)
        self.assertTrue(any(item[0].id == turn.id for item in pending))
        self.assertTrue(all(item[0].status == "pending" for item in pending))


if __name__ == "__main__":
    unittest.main()
