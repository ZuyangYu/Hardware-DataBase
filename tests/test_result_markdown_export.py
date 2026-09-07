"""Markdown answer structure must survive every result-export format."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from src.result_exports.content import envelope_from_turn
from src.result_exports.models import ResultEnvelope
from src.result_exports.renderers import render_result
from src.result_exports.markdown import parse_markdown_blocks


MARKDOWN_ANSWER = """## 文档基本信息对比

两份文档的项目对象不同，但文档类型一致。

| 对比项 | 知识库 ADAS HSI | 附件 EPS HSI | 结论 |
| --- | --- | --- | --- |
| 项目对象 | EQ6L T1G ADAS | EPS | 面向不同产品系统 |
| 发布日期 | 2025-02-24 | 2025-03-12 | EPS 较晚 |

### 结论

- ADAS 文档覆盖系统级架构。
- EPS 文档更关注 MCU 和 CAN 接口。

> 日期字段存在排版歧义，需人工核对。

```text
source_scope: attachment_and_knowledge_base
```
"""


HTML_FALLBACK_ANSWER = """## 对比结论

两份文档的项目对象不同。

## 二、PDF 报告生成

当前环境无法直接输出 PDF 文件，我已为您准备好自包含 HTML 报告。请按以下步骤操作即可获得 PDF：

将下方代码完整复制，保存为报告.html
用浏览器打开后打印为 PDF

```html
<!DOCTYPE html>
<html lang="zh-CN"><head><title>报告</title></head>
<body><h1>不应进入导出文件的 HTML 源码</h1></body></html>
```

补充说明：以上结论基于当前检索到的文本证据。
"""


def _markdown_envelope() -> ResultEnvelope:
    return ResultEnvelope(
        title="HSI 文档对比报告",
        query="对比知识库和附件中的 HSI 文档",
        answer=MARKDOWN_ANSWER,
        footer="已检索知识库和会话附件。",
        citations=[{"index": 1, "title": "600605919_EPS_HSI.docx", "locator": "第 1 页"}],
    )


def test_turn_envelope_persists_the_answer_block_model_for_future_renderers():
    turn = type(
        "Turn",
        (),
        {
            "query": "对比知识库和附件中的 HSI 文档",
            "answer": MARKDOWN_ANSWER,
            "footer": "",
            "summary": {},
            "kb_name": "hardware",
            "session_id": 3,
            "id": "turn-markdown-blocks",
            "query_mode": "deep",
        },
    )()

    envelope = envelope_from_turn(turn)

    assert envelope.blocks[0]["type"] == "heading"
    assert any(block["type"] == "table" for block in envelope.blocks)


def test_parse_markdown_answer_into_ordered_structured_blocks():
    blocks = parse_markdown_blocks(MARKDOWN_ANSWER)

    assert [block["type"] for block in blocks] == [
        "heading",
        "paragraph",
        "table",
        "heading",
        "list",
        "quote",
        "code",
    ]
    table = blocks[2]
    assert table["columns"] == ["对比项", "知识库 ADAS HSI", "附件 EPS HSI", "结论"]
    assert table["rows"][1] == ["发布日期", "2025-02-24", "2025-03-12", "EPS 较晚"]
    assert blocks[4]["ordered"] is False
    assert blocks[6]["language"] == "text"


def test_export_answer_removes_model_html_fallback_and_adds_async_export_status():
    import src.result_exports as exports

    normalize_export_answer = getattr(exports, "normalize_export_answer", lambda answer, _plan: answer)
    clean = normalize_export_answer(
        HTML_FALLBACK_ANSWER,
        {"formats": ["docx", "pdf"]},
    )

    assert "当前环境无法直接输出 PDF" not in clean
    assert "<!DOCTYPE html>" not in clean
    assert "<html" not in clean
    assert "两份文档的项目对象不同" in clean
    assert "补充说明" in clean
    assert "已提交 Word、PDF 导出任务，生成完成后可在下方下载。" in clean


def test_export_answer_normalization_is_idempotent_and_non_export_content_is_unchanged():
    import src.result_exports as exports

    normalize_export_answer = getattr(exports, "normalize_export_answer", lambda answer, _plan: answer)
    clean = normalize_export_answer("结论\n\n已提交 PDF 导出任务，生成完成后可在下方下载。", {"formats": ["pdf"]})

    assert normalize_export_answer(clean, {"formats": ["pdf"]}) == clean
    html_example = "```html\n<div>普通代码示例</div>\n```"
    assert normalize_export_answer(html_example, None) == html_example


def test_renderer_fallback_cleanup_keeps_an_ordinary_html_snippet():
    from src.result_exports.answer import strip_export_fallback_markup

    snippet = "示例：<html><p>这是正文片段</p></html>"

    assert strip_export_fallback_markup(snippet) == snippet


def test_docx_export_turns_markdown_table_into_a_word_table():
    rendered = render_result(_markdown_envelope(), "docx")
    document = Document(io.BytesIO(rendered.content))

    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows][1] == [
        "项目对象",
        "EQ6L T1G ADAS",
        "EPS",
        "面向不同产品系统",
    ]
    assert any(paragraph.text == "结论" for paragraph in document.paragraphs)
    assert all("| 对比项 |" not in paragraph.text for paragraph in document.paragraphs)


def test_pdf_export_turns_markdown_table_into_a_pdf_table():
    rendered = render_result(_markdown_envelope(), "pdf")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(rendered.content)).pages)

    assert b"/FontFile2" in rendered.content
    assert "项目对象" in text
    assert "EQ6L T1G ADAS" in text
    assert "面向不同产品系统" in text
    assert "| 对比项 |" not in text


def test_pdf_export_drops_a_full_html_document_fallback_from_old_answer():
    envelope = ResultEnvelope(
        title="HSI 文档对比报告",
        query="整理成 PDF",
        answer=HTML_FALLBACK_ANSWER,
        blocks=parse_markdown_blocks(HTML_FALLBACK_ANSWER),
    )
    rendered = render_result(envelope, "pdf")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(rendered.content)).pages)

    assert "不应进入导出文件的 HTML 源码" not in text
    assert "两份文档的项目对象不同" in text


def test_markdown_export_drops_a_full_html_document_fallback_from_old_snapshot():
    envelope = ResultEnvelope(
        title="HSI 文档对比报告",
        query="整理成 Markdown",
        answer=HTML_FALLBACK_ANSWER,
        blocks=parse_markdown_blocks(HTML_FALLBACK_ANSWER),
    )
    rendered = render_result(envelope, "md")
    text = rendered.content.decode("utf-8")

    assert "<!DOCTYPE html>" not in text
    assert "不应进入导出文件的 HTML 源码" not in text
    assert "两份文档的项目对象不同" in text


def test_xlsx_export_places_markdown_table_in_structured_cells():
    rendered = render_result(_markdown_envelope(), "xlsx")

    with zipfile.ZipFile(io.BytesIO(rendered.content)) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
        sheets = [
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet")
        ]

    assert "Markdown 表格 1" in workbook
    assert any("EQ6L T1G ADAS" in sheet and "面向不同产品系统" in sheet for sheet in sheets)
    assert not any("| 对比项 |" in sheet for sheet in sheets)


def test_pptx_export_places_markdown_table_in_a_table_shape():
    rendered = render_result(_markdown_envelope(), "pptx")
    presentation = Presentation(io.BytesIO(rendered.content))

    table_shapes = [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_table
    ]
    assert table_shapes
    cells = [cell.text for row in table_shapes[0].table.rows for cell in row.cells]
    assert "项目对象" in cells
    assert "EQ6L T1G ADAS" in cells
    assert "| 对比项 |" not in cells
