from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from pypdf import PdfReader

from src.document_authoring.conversion import (
    TemplateArtifactConversionService,
    TemplateConversionError,
)
from src.document_authoring.models import DocumentArtifact
from src.document_authoring.service import DocumentGenerationService


def _docx_fixture() -> bytes:
    document = Document()
    document.add_heading("600608964 ADAS ICD", level=1)
    document.add_paragraph("工作电压范围：9 ~ 16V")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "管脚"
    table.rows[0].cells[1].text = "功能"
    row = table.add_row().cells
    row[0].text = "1"
    row[1].text = "UBD"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_fixture() -> bytes:
    content_types = """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""
    workbook = """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="ICD" sheetId="1" r:id="rId1"/></sheets></workbook>"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>"""
    sheet = """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>管脚</t></is></c><c r="B1" t="inlineStr"><is><t>功能</t></is></c></row>
<row r="2"><c r="A2" t="inlineStr"><is><t>1</t></is></c><c r="B2" t="inlineStr"><is><t>UBD</t></is></c></row></sheetData></worksheet>"""
    output = BytesIO()
    with ZipFile(output, "w") as package:
        package.writestr("[Content_Types].xml", content_types)
        package.writestr("_rels/.rels", rels)
        package.writestr("xl/workbook.xml", workbook)
        package.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        package.writestr("xl/worksheets/sheet1.xml", sheet)
    return output.getvalue()


def test_docx_to_pdf_is_semantic_and_validated() -> None:
    result = TemplateArtifactConversionService().convert(
        _docx_fixture(), source_format="docx", target_format="pdf"
    )

    assert result.content.startswith(b"%PDF")
    assert result.metadata["source_format"] == "docx"
    assert result.metadata["target_format"] == "pdf"
    assert result.metadata["page_count"] >= 1
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(result.content)).pages)
    assert "600608964" in text
    assert "UBD" in text


def test_xlsx_to_pdf_preserves_sheet_rows() -> None:
    result = TemplateArtifactConversionService().convert(
        _xlsx_fixture(), source_format="xlsx", target_format="pdf"
    )

    assert result.content.startswith(b"%PDF")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(result.content)).pages)
    assert "管脚" in text
    assert "UBD" in text
    assert result.metadata["sheet_count"] == 1


@pytest.mark.parametrize("source_format, fixture", [("docx", _docx_fixture), ("xlsx", _xlsx_fixture)])
def test_native_office_to_pptx_is_a_valid_package(source_format: str, fixture) -> None:
    result = TemplateArtifactConversionService().convert(
        fixture(), source_format=source_format, target_format="pptx"
    )

    assert result.content[:2] == b"PK"
    with ZipFile(BytesIO(result.content)) as package:
        names = set(package.namelist())
        assert "[Content_Types].xml" in names
        assert "ppt/presentation.xml" in names
        assert any(name.startswith("ppt/slides/slide") for name in names)
    assert result.metadata["slide_count"] >= 1


def test_conversion_rejects_unsupported_or_active_content() -> None:
    with pytest.raises(TemplateConversionError, match="unsupported target"):
        TemplateArtifactConversionService().convert(
            _docx_fixture(), source_format="docx", target_format="html"
        )

    active = BytesIO()
    with ZipFile(active, "w") as package:
        package.writestr("[Content_Types].xml", "<Types/>")
        package.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(TemplateConversionError, match="active content"):
        TemplateArtifactConversionService().convert(
            active.getvalue(), source_format="xlsm", target_format="pdf"
        )


def test_document_service_persists_converted_child_artifact() -> None:
    source_content = _docx_fixture()
    source = DocumentArtifact(
        artifact_id="artifact-native",
        work_order_id="wo-1",
        run_id="run-1",
        output_format="docx",
        stage="review_candidate",
        content_hash=__import__("hashlib").sha256(source_content).hexdigest(),
        validation_report_id="report-native",
        integrity_manifest_id="manifest-native",
    )
    saved: list[tuple[DocumentArtifact, bytes, str]] = []
    store = SimpleNamespace(
        get_artifact=lambda artifact_id: source if artifact_id == source.artifact_id else None,
        read_artifact_content=lambda _artifact_id: source_content,
        save_validation_report=Mock(),
        save_artifact=lambda artifact, content, suffix: (saved.append((artifact, content, suffix)) or artifact),
    )
    service = object.__new__(DocumentGenerationService)
    service.store = store
    service.converter = TemplateArtifactConversionService()
    service._artifact_for_context = lambda _ctx, _artifact_id: source
    service._order = Mock(return_value=SimpleNamespace(
        work_order_id="wo-1",
        target_format="docx",
        scope_type="project",
    ))
    converted = service.convert_document_artifact(SimpleNamespace(), "artifact-native", target_format="pdf")

    assert converted.output_format == "pdf"
    assert converted.parent_artifact_id == "artifact-native"
    assert saved and saved[0][2] == "pdf"
    assert saved[0][1].startswith(b"%PDF")
