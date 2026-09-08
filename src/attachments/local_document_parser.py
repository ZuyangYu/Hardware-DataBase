"""Deterministic local parsers for chat attachments (design §7).

PDF (pypdf)  : page -> text blocks -> parts; low text density degrades the
               asset (``ocr_candidate``) instead of fabricating content.
DOCX         : python-docx headings/paragraphs/tables with heading-path
               locators; no fake page numbers.
TXT/MD       : bounded paragraph chunks with line locators.

All parsers enforce hard resource limits from settings (page caps, cell caps,
uncompressed size caps) *during* parsing, because checking after the fact is
too late for zip bombs.
"""

from __future__ import annotations

import hashlib
import json
import os
import re

import src.settings
from src.attachments.models import (
    AttachmentAsset,
    AttachmentPart,
    PARSE_STATUS_DEGRADED,
    PARSE_STATUS_FAILED,
    PARSE_STATUS_READY,
    PART_TYPE_IMAGE,
    PART_TYPE_OCR_TEXT,
    PART_TYPE_TABLE,
    PART_TYPE_TEXT,
    PARSER_VERSION,
)
from src.attachments.ocr import OcrEngine, default_ocr_engine
from src.attachments.router import ParseOutcome

# Average chars per token used for budget math only (never shown as truth).
_CHARS_PER_TOKEN = 3.5
_MAX_PART_TEXT_CHARS = 8000
_MIN_PAGE_TEXT_DENSITY = 24  # chars; below this a PDF page is "image-only"


def _content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", "replace")).hexdigest()


def _part(asset_id: str, ordinal: int, part_type: str, text: str, locator: dict, metadata: dict | None = None) -> AttachmentPart:
    return AttachmentPart(
        part_id="",
        asset_id=asset_id,
        ordinal=ordinal,
        part_type=part_type,
        text_content=text[:_MAX_PART_TEXT_CHARS],
        locator=locator,
        metadata=metadata or {},
        content_hash=_content_hash(text),
        parser_version=PARSER_VERSION,
    )


def _rough_tokens(text: str) -> int:
    return int(len(text) / _CHARS_PER_TOKEN)


def _json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _split_oversized_docx_row(row: list[str], max_cell_text_length: int) -> list[list[str]]:
    """Represent an oversized source row as ordered, bounded cell fragments.

    A fragment keeps the original cell position and leaves the other cells
    empty.  This keeps every payload a JSON array of rows while allowing the
    caller to recover the exact source cell text by concatenating fragments.
    """
    fragments: list[list[str]] = []
    width = len(row)
    for cell_index, value in enumerate(row):
        cursor = 0
        if not value:
            values = [""]
        else:
            values = []
            while cursor < len(value):
                remaining = min(max_cell_text_length, len(value) - cursor)
                low, high = 1, remaining
                best = 1
                while low <= high:
                    middle = (low + high) // 2
                    candidate = [""] * width
                    candidate[cell_index] = value[cursor : cursor + middle]
                    if len(_json_text([candidate])) <= _MAX_PART_TEXT_CHARS:
                        best = middle
                        low = middle + 1
                    else:
                        high = middle - 1
                values.append(value[cursor : cursor + best])
                cursor += best
        for fragment in values:
            candidate = [""] * width
            candidate[cell_index] = fragment
            fragments.append(candidate)
    return fragments


def _docx_table_chunks(
    rows: list[list[str]], max_cell_text_length: int
) -> list[tuple[str, int, int]]:
    """Serialize DOCX rows into valid, bounded JSON payloads.

    The row bounds are source row indices (zero based), not chunk-local
    positions.  Oversized rows are split into cell fragments before the
    normal row packing loop, so no ``_part`` truncation can corrupt JSON.
    """
    chunks: list[tuple[str, int, int]] = []
    current: list[list[str]] = []
    current_indices: list[int] = []

    def flush() -> None:
        if not current:
            return
        chunks.append((_json_text(current), min(current_indices), max(current_indices)))
        current.clear()
        current_indices.clear()

    for row_index, row in enumerate(rows):
        if len(_json_text([row])) <= _MAX_PART_TEXT_CHARS and all(
            len(cell) <= max_cell_text_length for cell in row
        ):
            row_fragments = [row]
        else:
            row_fragments = _split_oversized_docx_row(row, max_cell_text_length)
        for row_fragment in row_fragments:
            candidate = current + [row_fragment]
            if current and len(_json_text(candidate)) > _MAX_PART_TEXT_CHARS:
                flush()
                candidate = [row_fragment]
            # _split_oversized_docx_row guarantees this branch for ordinary
            # strings; retain a defensive fallback for unusual JSON overhead.
            if len(_json_text(candidate)) > _MAX_PART_TEXT_CHARS:
                continue
            current.append(row_fragment)
            current_indices.append(row_index)
    flush()
    return chunks


class LocalDocumentParser:
    """PDF / DOCX / TXT / MD parsing for the attachment domain."""

    kind = "local_document"

    def __init__(self, ocr_engine: OcrEngine | None = None):
        self._ocr_engine = ocr_engine

    def parse(self, asset: AttachmentAsset, source_path: str) -> ParseOutcome:
        ext = (asset.extension or "").lower()
        try:
            if ext == ".pdf":
                return self._parse_pdf(asset, source_path)
            if ext == ".docx":
                return self._parse_docx(asset, source_path)
            if ext in {".txt", ".md"}:
                return self._parse_plaintext(asset, source_path, ext)
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={},
                error_code="unsupported_document_type",
                error_message=f"local document parser cannot handle {ext}",
            )
        except MemoryError:
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={},
                error_code="parse_resource_exhausted",
                error_message="document too large to parse locally",
            )
        except Exception as exc:  # noqa: BLE001 - worker converts to degraded/failed
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={},
                error_code="parse_failed",
                error_message=str(exc)[:500],
            )

    # -- PDF ----------------------------------------------------------------

    def _parse_pdf(self, asset: AttachmentAsset, source_path: str) -> ParseOutcome:
        from pypdf import PdfReader

        max_pages = max(1, int(src.settings.CHAT_ATTACHMENT_MAX_PAGES))
        reader = PdfReader(source_path)
        total_pages = len(reader.pages)
        if total_pages > max_pages:
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={"page_count": total_pages},
                error_code="page_limit_exceeded",
                error_message=f"PDF has {total_pages} pages; limit is {max_pages}",
            )
        parts: list[AttachmentPart] = []
        degraded_reasons: list[str] = []
        low_density_pages: list[int] = []
        ocr_attempted_pages: list[int] = []
        ocr_indexed_pages: list[int] = []
        ocr_failed_pages: list[int] = []
        empty_pages = 0
        ordinal = 0
        for page_index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            text = text.strip()
            if len(text) < _MIN_PAGE_TEXT_DENSITY:
                empty_pages += 1
                low_density_pages.append(page_index)
                continue
            blocks = _split_text_blocks(text)
            for block_index, block in enumerate(blocks, start=1):
                parts.append(
                    _part(
                        asset.asset_id,
                        ordinal,
                        PART_TYPE_TEXT,
                        block,
                        locator={"page": page_index, "block": block_index},
                        metadata={"format": "pdf"},
                    )
                )
                ordinal += 1
        if empty_pages and src.settings.CHAT_ATTACHMENT_OCR_ENABLED:
            max_ocr_pages = max(0, int(src.settings.CHAT_ATTACHMENT_OCR_MAX_PAGES))
            candidate_pages = low_density_pages[:max_ocr_pages]
            engine = self._ocr_engine or default_ocr_engine()
            provider = str(getattr(engine, "provider", "local_cpu"))
            for page_index in candidate_pages:
                ocr_attempted_pages.append(page_index)
                try:
                    ocr_text = str(engine.recognize_page(source_path, page_index) or "").strip()
                except Exception as exc:  # noqa: BLE001 - one page must not abort the asset
                    ocr_failed_pages.append(page_index)
                    detail = str(exc).strip()[:160] or "OCR failed"
                    degraded_reasons.append(f"OCR page {page_index} failed: {detail}")
                    continue
                if not ocr_text:
                    ocr_failed_pages.append(page_index)
                    degraded_reasons.append(f"OCR page {page_index} produced no text")
                    continue
                ocr_blocks = _split_text_blocks(ocr_text)
                if not ocr_blocks:
                    ocr_failed_pages.append(page_index)
                    degraded_reasons.append(f"OCR page {page_index} produced no text")
                    continue
                for block_index, block in enumerate(ocr_blocks, start=1):
                    parts.append(
                        _part(
                            asset.asset_id,
                            ordinal,
                            PART_TYPE_OCR_TEXT,
                            block,
                            locator={"page": page_index, "block": block_index},
                            metadata={"format": "pdf", "provider": provider},
                        )
                    )
                    ordinal += 1
                ocr_indexed_pages.append(page_index)

            unresolved_pages = [
                page for page in low_density_pages if page not in ocr_indexed_pages
            ]
            if unresolved_pages:
                degraded_reasons.insert(
                    0,
                    f"{len(unresolved_pages)} page(s) contain no extractable text "
                    "(likely scanned or OCR incomplete)",
                )
            try:
                from src.observability.metrics import record_attachment

                record_attachment(
                    "ocr",
                    status="ok" if ocr_indexed_pages else "degraded",
                )
            except Exception:
                pass
        elif empty_pages:
            unresolved_pages = list(low_density_pages)
            degraded_reasons.append(
                f"{empty_pages} page(s) contain no extractable text (likely scanned)"
            )
        else:
            unresolved_pages = []
        manifest = {
            "format": "pdf",
            "page_count": total_pages,
            "indexed_pages": total_pages - empty_pages,
            "low_text_density_pages": low_density_pages[:50],
            "ocr_candidate": bool(unresolved_pages),
            "ocr_enabled": bool(src.settings.CHAT_ATTACHMENT_OCR_ENABLED),
            "ocr_attempted_pages": ocr_attempted_pages,
            "ocr_indexed_pages": ocr_indexed_pages,
            "ocr_failed_pages": ocr_failed_pages,
            "part_count": len(parts),
            "rough_token_count": sum(_rough_tokens(p.text_content) for p in parts),
        }
        status = PARSE_STATUS_DEGRADED if degraded_reasons else PARSE_STATUS_READY
        return ParseOutcome(
            ok=True,
            parse_status=status,
            parts=parts,
            manifest=manifest,
            degraded_reasons=degraded_reasons,
        )

    # -- DOCX ---------------------------------------------------------------

    def _parse_docx(self, asset: AttachmentAsset, source_path: str) -> ParseOutcome:
        import docx  # python-docx

        _check_ooxml_limits(source_path)
        document = docx.Document(source_path)
        parts: list[AttachmentPart] = []
        degraded_reasons: list[str] = []
        heading_path: list[str] = []
        ordinal = 0
        table_index = 0
        table_part_count = 0
        degraded_table_count = 0
        rows_remaining = max(0, int(src.settings.CHAT_ATTACHMENT_MAX_ROWS))
        cells_remaining = max(0, int(src.settings.CHAT_ATTACHMENT_MAX_CELLS))

        def _text_of_cell(cell) -> str:
            return "\n".join(
                paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()
            )

        # document.iter_inner_content keeps paragraphs and tables in document
        # order on python-docx >= 1.1.
        try:
            flow_items = list(document.iter_inner_content())
        except AttributeError:  # very old python-docx
            flow_items = list(document.paragraphs)
        for item in flow_items:
            style_name = str(getattr(item, "style", None) and item.style.name or "")
            if hasattr(item, "tables"):  # Document container guard
                continue
            if getattr(item, "text", None) is None and not hasattr(item, "rows"):
                continue
            if hasattr(item, "rows"):  # Table
                source_rows = item.rows
                max_cell_text_length = max(
                    1, int(src.settings.CHAT_ATTACHMENT_MAX_CELL_TEXT_LENGTH)
                )
                rows: list[list[str]] = []
                table_degraded = False
                allowed_rows = min(len(source_rows), rows_remaining)
                if len(source_rows) > allowed_rows:
                    degraded_reasons.append(
                        f"DOCX table {table_index} exceeds row limit; "
                        f"{len(source_rows) - allowed_rows} row(s) not indexed"
                    )
                    table_degraded = True
                cells_seen = 0
                fragmented_row_indices: set[int] = set()
                for row_index, row in enumerate(source_rows[:allowed_rows]):
                    remaining_cells = cells_remaining - cells_seen
                    if remaining_cells <= 0:
                        degraded_reasons.append(
                            f"DOCX table {table_index} exceeds cell limit; "
                            "remaining rows not indexed"
                        )
                        table_degraded = True
                        break
                    cells = row.cells
                    if len(cells) > remaining_cells:
                        degraded_reasons.append(
                            f"DOCX table {table_index} exceeds cell limit; "
                            f"{len(cells) - remaining_cells} cell(s) in row {row_index} not indexed"
                        )
                        table_degraded = True
                    selected_cells = cells[:remaining_cells]
                    rows.append([_text_of_cell(cell) for cell in selected_cells])
                    cells_seen += len(selected_cells)
                    if selected_cells and (
                        len(_json_text([rows[-1]])) > _MAX_PART_TEXT_CHARS
                        or any(
                            len(value) > max_cell_text_length for value in rows[-1]
                        )
                    ):
                        fragmented_row_indices.add(row_index)
                rows_remaining -= len(rows)
                cells_remaining -= cells_seen
                if not rows:
                    if table_degraded:
                        degraded_table_count += 1
                    table_index += 1
                    continue
                header_context = list(rows[0])
                table_locator = ordinal
                for table_text, row_start, row_end in _docx_table_chunks(
                    rows, max_cell_text_length
                ):
                    payload_rows = json.loads(table_text)
                    parts.append(
                        _part(
                            asset.asset_id,
                            ordinal,
                            PART_TYPE_TABLE,
                            table_text,
                            locator={
                                "heading_path": list(heading_path),
                                # Preserve the historical table locator (the
                                # first emitted part's ordinal); table_index
                                # is the stable source-table identity.
                                "table": table_locator,
                                "table_index": table_index,
                                "row_start": row_start,
                                "row_end": row_end,
                            },
                            metadata={
                                "format": "docx",
                                "rows": len(payload_rows),
                                "table_rows": len(rows),
                                "table_index": table_index,
                                "table_id": f"table-{table_index}",
                                "row_start": row_start,
                                "row_end": row_end,
                                "header_context": header_context,
                                "row_fragment": row_start == row_end
                                and row_start in fragmented_row_indices,
                            },
                        )
                    )
                    ordinal += 1
                    table_part_count += 1
                if table_degraded:
                    degraded_table_count += 1
                table_index += 1
                continue
            text = (item.text or "").strip()
            if not text:
                continue
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
                heading_path = heading_path[: max(0, level - 1)]
                heading_path.append(text[:200])
            parts.append(
                _part(
                    asset.asset_id,
                    ordinal,
                    PART_TYPE_TEXT,
                    text,
                    locator={"heading_path": list(heading_path), "paragraph": ordinal},
                    metadata={"format": "docx", "style": style_name},
                )
            )
            ordinal += 1

        # Relationships are metadata only: external hyperlinks are never
        # fetched and embedded image bytes are never executed. Keeping them as
        # canonical parts lets retrieval/document flow cite their presence
        # without inventing page coordinates.
        links: list[dict[str, str]] = []
        embedded_images: list[dict[str, str]] = []
        for relationship_id, relationship in document.part.rels.items():
            reltype = str(getattr(relationship, "reltype", "") or "")
            if reltype.endswith("/hyperlink"):
                target = str(getattr(relationship, "target_ref", "") or "").strip()
                if target:
                    links.append({"relationship_id": str(relationship_id), "target": target})
                    parts.append(
                        _part(
                            asset.asset_id,
                            ordinal,
                            PART_TYPE_TEXT,
                            f"链接: {target}",
                            locator={"relationship_id": str(relationship_id), "kind": "link"},
                            metadata={"format": "docx", "kind": "link", "target": target},
                        )
                    )
                    ordinal += 1
            elif reltype.endswith("/image"):
                target_part = getattr(relationship, "target_part", None)
                image = {
                    "relationship_id": str(relationship_id),
                    "part_name": str(getattr(target_part, "partname", "") or ""),
                    "content_type": str(getattr(target_part, "content_type", "") or ""),
                }
                embedded_images.append(image)
                parts.append(
                    _part(
                        asset.asset_id,
                        ordinal,
                        PART_TYPE_IMAGE,
                        "",
                        locator={"relationship_id": str(relationship_id), "kind": "embedded_image"},
                        metadata={"format": "docx", "kind": "embedded_image", **image},
                    )
                )
                ordinal += 1

        if not parts:
            degraded_reasons.append("document contains no extractable paragraphs or tables")
        manifest = {
            "format": "docx",
            "part_count": len(parts),
            "rough_token_count": sum(_rough_tokens(p.text_content) for p in parts),
            "link_count": len(links),
            "links": links[:100],
            "embedded_image_count": len(embedded_images),
            "embedded_images": embedded_images[:100],
            "table_count": table_index,
            "table_part_count": table_part_count,
            "degraded_table_count": degraded_table_count,
        }
        status = PARSE_STATUS_DEGRADED if degraded_reasons else PARSE_STATUS_READY
        return ParseOutcome(
            ok=True,
            parse_status=status,
            parts=parts,
            manifest=manifest,
            degraded_reasons=degraded_reasons,
        )

    # -- TXT / MD -----------------------------------------------------------

    def _parse_plaintext(self, asset: AttachmentAsset, source_path: str, ext: str) -> ParseOutcome:
        max_bytes = int(src.settings.CHAT_ATTACHMENT_MAX_UNCOMPRESSED_BYTES)
        size = os.path.getsize(source_path)
        if size > max_bytes:
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={},
                error_code="file_limit_exceeded",
                error_message=f"file exceeds {max_bytes} bytes",
            )
        try:
            text = open(source_path, "r", encoding="utf-8", errors="replace").read()
        except OSError as exc:
            return ParseOutcome(
                ok=False,
                parse_status=PARSE_STATUS_FAILED,
                parts=[],
                manifest={},
                error_code="read_failed",
                error_message=str(exc)[:300],
            )
        lines = text.splitlines()
        parts: list[AttachmentPart] = []
        buffer: list[str] = []
        buffer_start_line = 1
        ordinal = 0
        line_cursor = 0
        for line in lines:
            line_cursor += 1
            buffer.append(line)
            joined = "\n".join(buffer)
            if len(joined) >= 1800 or (line.strip() == "" and joined.strip()):
                if joined.strip():
                    parts.append(
                        _part(
                            asset.asset_id,
                            ordinal,
                            PART_TYPE_TEXT,
                            joined.strip(),
                            locator={"line_start": buffer_start_line, "line_end": line_cursor},
                            metadata={"format": ext.lstrip(".")},
                        )
                    )
                    ordinal += 1
                buffer = []
                buffer_start_line = line_cursor + 1
        tail = "\n".join(buffer).strip()
        if tail:
            parts.append(
                _part(
                    asset.asset_id,
                    ordinal,
                    PART_TYPE_TEXT,
                    tail,
                    locator={"line_start": buffer_start_line, "line_end": line_cursor},
                    metadata={"format": ext.lstrip(".")},
                )
            )
            ordinal += 1
        manifest = {
            "format": ext.lstrip("."),
            "line_count": line_cursor,
            "part_count": len(parts),
            "rough_token_count": sum(_rough_tokens(p.text_content) for p in parts),
        }
        return ParseOutcome(
            ok=True,
            parse_status=PARSE_STATUS_READY,
            parts=parts,
            manifest=manifest,
        )


def _split_text_blocks(text: str) -> list[str]:
    """Split a page's text into stable, size-bounded blocks."""
    raw_blocks = re.split(r"\n\s*\n", text)
    blocks: list[str] = []
    buffer = ""
    for block in raw_blocks:
        block = block.strip()
        if not block:
            continue
        if len(buffer) + len(block) + 1 <= 1200:
            buffer = f"{buffer}\n{block}".strip()
        else:
            if buffer:
                blocks.append(buffer)
            if len(block) <= 1200:
                buffer = block
            else:
                # Oversized single block: chunk on line boundaries.
                current = ""
                for line in block.splitlines():
                    if len(current) + len(line) + 1 > 1200 and current:
                        blocks.append(current)
                        current = line
                    else:
                        current = f"{current}\n{line}".strip()
                buffer = current
    if buffer:
        blocks.append(buffer)
    return blocks


def _check_ooxml_limits(source_path: str) -> None:
    """Enforce zip-bomb style limits *before* python-docx reads the package."""
    import zipfile

    max_uncompressed = int(src.settings.CHAT_ATTACHMENT_MAX_UNCOMPRESSED_BYTES)
    max_ratio = 200
    with zipfile.ZipFile(source_path) as bundle:
        total_uncompressed = 0
        for info in bundle.infolist():
            total_uncompressed += info.file_size
            if total_uncompressed > max_uncompressed:
                raise ValueError("OOXML package exceeds the uncompressed size limit")
            if info.compress_size and info.file_size / max(1, info.compress_size) > max_ratio:
                raise ValueError("OOXML package has a suspicious compression ratio")
