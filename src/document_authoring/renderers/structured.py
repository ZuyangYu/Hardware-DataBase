"""Deterministic renderers for server-owned template-free recipes.

These renderers create fresh packages from a validated :class:`DocumentModel`.
They do not accept template bytes, coordinates, XML snippets, macros, external
relationships, or caller-selected styles.  The recipe binding is the only
layout input and is produced by ``StructureBindingCompiler``.
"""

from __future__ import annotations

import copy
import hashlib
import io
import math
import zipfile
from dataclasses import dataclass
from datetime import datetime
from typing import Any
from xml.sax.saxutils import escape

from src.document_authoring.document_model import (
    DocumentModel,
    ListBlock,
    ParagraphBlock,
    SectionBlock,
    TypedTableBlock,
)
from src.document_authoring.models import TypedTableRow, content_hash
from src.document_authoring.ooxml import sanitize_xml10_text, validate_ooxml_package
from src.document_authoring.planning.recipes import (
    StructureBindingSet,
    StructureConstraintsProfile,
    SystemRecipe,
    selected_profile,
)


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"
_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class StructuredRenderResult:
    """One safe fresh-package rendering and its hash-bound manifest."""

    content: bytes
    mime_type: str
    extension: str
    preview: dict[str, Any]
    integrity_manifest: dict[str, Any]


class _StructuredRenderer:
    format: str = ""
    mime_type: str = ""

    def _validate(
        self,
        model: DocumentModel,
        recipe: SystemRecipe,
        binding: StructureBindingSet,
    ) -> StructureConstraintsProfile:
        if not isinstance(model, DocumentModel):
            model = DocumentModel.model_validate(model)
        if recipe.status != "available":
            raise ValueError("cannot render with an unavailable system recipe")
        if self.format not in recipe.supported_formats:
            raise ValueError(f"recipe does not support output format: {self.format}")
        if binding.output_format != self.format:
            raise ValueError("structure binding output format does not match renderer")
        if binding.recipe_id != recipe.recipe_id or binding.recipe_version != recipe.version:
            raise ValueError("structure binding is not for the selected recipe")
        if binding.recipe_hash != (recipe.recipe_hash or ""):
            raise ValueError("structure binding recipe hash does not match the selected recipe")
        if model.plan_id and model.plan_id != binding.plan_id:
            raise ValueError("document model plan identity does not match structure binding")
        if model.plan_hash and model.plan_hash != binding.plan_hash:
            raise ValueError("document model plan hash does not match structure binding")
        if len(model.blocks) != len(binding.bindings):
            raise ValueError("structure binding does not cover every document block")
        by_unit = {item.unit_id: item for item in binding.bindings}
        if len(by_unit) != len(binding.bindings):
            raise ValueError("structure binding contains duplicate units")
        for ordinal, block in enumerate(model.blocks):
            item = by_unit.get(block.unit_id)
            if item is None or item.ordinal != ordinal or item.block_id != block.block_id:
                raise ValueError("structure binding order or block identity does not match model")
            if item.block_kind != block.kind:
                raise ValueError("structure binding block kind does not match model")
            if item.component_id not in recipe.component_ids:
                raise ValueError("structure binding references an unregistered component")
            hints = getattr(block, "layout_hints", {}) or {}
            unsafe = set(hints) - {"level", "emphasis", "orientation", "keep_with_next"}
            if unsafe:
                raise ValueError(f"layout hint is not allowlisted: {sorted(unsafe)}")
        profile = selected_profile(recipe)
        if len(model.blocks) > profile.max_components:
            raise ValueError("structure exceeds the recipe component limit")
        text_size = 0
        for block in model.blocks:
            block_text, rows, columns = _block_size(block)
            text_size += block_text
            if rows > profile.max_table_rows:
                raise ValueError("structure exceeds the recipe table row limit")
            if columns > profile.max_table_columns:
                raise ValueError("structure exceeds the recipe table column limit")
        if text_size > profile.max_text_chars:
            raise ValueError("structure exceeds the recipe text limit")
        return profile

    def _result(
        self,
        content: bytes,
        *,
        model: DocumentModel,
        recipe: SystemRecipe,
        binding: StructureBindingSet,
        profile: StructureConstraintsProfile,
        preview: dict[str, Any],
    ) -> StructuredRenderResult:
        if len(content) > profile.max_artifact_bytes:
            raise ValueError("generated artifact exceeds the recipe byte limit")
        artifact_hash = hashlib.sha256(content).hexdigest()
        manifest = {
            "format": self.format,
            "recipe_id": recipe.recipe_id,
            "recipe_version": recipe.version,
            "recipe_hash": recipe.recipe_hash,
            "profile_id": profile.profile_id,
            "profile_version": profile.version,
            "plan_id": binding.plan_id,
            "plan_version": binding.plan_version,
            "plan_hash": binding.plan_hash,
            "binding_hash": binding.binding_hash,
            "document_model_hash": model.model_hash,
            "artifact_hash": artifact_hash,
            "active_content_status": "clean",
            "external_relationships": [],
            "macro_parts": [],
            "package_parts": [],
            "policy_violations": [],
        }
        manifest["manifest_hash"] = content_hash(manifest)
        return StructuredRenderResult(
            content=content,
            mime_type=self.mime_type,
            extension=self.format,
            preview={**preview, "format": self.format, "artifact_hash": artifact_hash},
            integrity_manifest=manifest,
        )


class StructuredDocxRenderer(_StructuredRenderer):
    """Create a clean DOCX from recipe-approved semantic components."""

    format = "docx"
    mime_type = DOCX_MIME

    def render(
        self,
        document_model: DocumentModel,
        recipe: SystemRecipe,
        binding: StructureBindingSet,
    ) -> StructuredRenderResult:
        profile = self._validate(document_model, recipe, binding)
        from docx import Document
        from docx.shared import Inches, Pt

        document = Document()
        document.core_properties.author = "Hardware DataBase"
        document.core_properties.last_modified_by = "Hardware DataBase"
        document.core_properties.created = datetime(2000, 1, 1)
        document.core_properties.modified = datetime(2000, 1, 1)
        document.core_properties.title = "Generated document"
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        block_count = 0
        table_count = 0
        for block in document_model.blocks:
            if isinstance(block, SectionBlock):
                title = sanitize_xml10_text(block.title or block.content or block.unit_id)
                document.add_heading(title, level=_heading_level(block))
                if block.content and block.title:
                    document.add_paragraph(sanitize_xml10_text(block.content))
                block_count += 1
            elif isinstance(block, ParagraphBlock):
                document.add_paragraph(sanitize_xml10_text(block.content))
                block_count += 1
            elif isinstance(block, ListBlock):
                for item in block.items:
                    document.add_paragraph(sanitize_xml10_text(item), style="List Bullet")
                block_count += 1
            elif isinstance(block, TypedTableBlock):
                document.add_table(
                    rows=1,
                    cols=len(block.columns) + 1,
                )
                table = document.tables[-1]
                table.style = "Table Grid"
                headers = ["row_key", *block.columns]
                for index, header in enumerate(headers):
                    table.rows[0].cells[index].text = sanitize_xml10_text(header)
                for row in block.rows:
                    cells = table.add_row().cells
                    cells[0].text = sanitize_xml10_text(row.row_key)
                    for index, column in enumerate(block.columns, start=1):
                        cells[index].text = sanitize_xml10_text(row.cells.get(column, ""))
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                block_count += 1
                table_count += 1
            else:
                # Cross-reference blocks have no free-form markup; a plain
                # paragraph is the only representation exposed by this recipe.
                references = getattr(block, "references", []) or []
                document.add_paragraph(sanitize_xml10_text(", ".join(references)))
                block_count += 1

        raw = io.BytesIO()
        document.save(raw)
        rendered = _canonical_zip(raw.getvalue())
        validate_ooxml_package(rendered, "docx")
        with zipfile.ZipFile(io.BytesIO(rendered), "r") as package:
            names = package.namelist()
            if any("vba" in name.lower() or "/embeddings/" in name.lower() for name in names):
                raise ValueError("template-free DOCX must not contain active content")
        return self._result(
            rendered,
            model=document_model,
            recipe=recipe,
            binding=binding,
            profile=profile,
            preview={"block_count": block_count, "table_count": table_count},
        )


class StructuredPdfRenderer(_StructuredRenderer):
    """Create a bounded PDF from the same semantic model and recipe."""

    format = "pdf"
    mime_type = PDF_MIME

    def render(
        self,
        document_model: DocumentModel,
        recipe: SystemRecipe,
        binding: StructureBindingSet,
    ) -> StructuredRenderResult:
        profile = self._validate(document_model, recipe, binding)
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "RecipeTitle", parent=styles["Title"], fontName="Helvetica",
            fontSize=18, leading=22, alignment=TA_CENTER, spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            "RecipeHeading", parent=styles["Heading2"], fontName="Helvetica",
            fontSize=13, leading=17, spaceBefore=8, spaceAfter=5,
        )
        body_style = ParagraphStyle(
            "RecipeBody", parent=styles["BodyText"], fontName="Helvetica",
            fontSize=10, leading=14, spaceAfter=6,
        )
        small_style = ParagraphStyle(
            "RecipeSmall", parent=body_style, fontSize=8, leading=10,
        )
        story: list[Any] = [Paragraph("Generated document", title_style)]
        block_count = 0
        table_count = 0
        for block in document_model.blocks:
            if isinstance(block, SectionBlock):
                story.append(Paragraph(_pdf_text(block.title or block.content or block.unit_id), heading_style))
                if block.content and block.title:
                    story.append(Paragraph(_pdf_text(block.content), body_style))
            elif isinstance(block, ParagraphBlock):
                story.append(Paragraph(_pdf_text(block.content), body_style))
            elif isinstance(block, ListBlock):
                for index, item in enumerate(block.items, start=1):
                    story.append(Paragraph(_pdf_text(f"{index}. {item}"), body_style))
            elif isinstance(block, TypedTableBlock):
                values = [[_pdf_text("row_key"), *(_pdf_text(column) for column in block.columns)]]
                values.extend([
                    [_pdf_text(row.row_key), *(_pdf_text(row.cells.get(column, "")) for column in block.columns)]
                    for row in block.rows
                ])
                table = Table(values, repeatRows=1, colWidths=[55] + [None] * len(block.columns))
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C3D4")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.extend([table, Spacer(1, 4 * mm)])
                table_count += 1
            else:
                references = getattr(block, "references", []) or []
                story.append(Paragraph(_pdf_text(", ".join(references)), body_style))
            block_count += 1

        def invariant_canvas(filename: Any, **kwargs: Any):
            kwargs["invariant"] = 1
            return canvas.Canvas(filename, **kwargs)

        output = io.BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=16 * mm,
            bottomMargin=16 * mm,
            title="Generated document",
            author="Hardware DataBase",
        )
        document.build(story, canvasmaker=invariant_canvas)
        rendered = output.getvalue()
        if not rendered.startswith(b"%PDF") or not rendered.rstrip().endswith(b"%%EOF"):
            raise ValueError("generated PDF is not a complete PDF package")
        return self._result(
            rendered,
            model=document_model,
            recipe=recipe,
            binding=binding,
            profile=profile,
            preview={"block_count": block_count, "table_count": table_count},
        )


class StructuredXlsxRenderer(_StructuredRenderer):
    """Create a clean XLSX workbook with only inline, non-formula values."""

    format = "xlsx"
    mime_type = XLSX_MIME

    def render(
        self,
        document_model: DocumentModel,
        recipe: SystemRecipe,
        binding: StructureBindingSet,
    ) -> StructuredRenderResult:
        profile = self._validate(document_model, recipe, binding)
        table_blocks = [block for block in document_model.blocks if isinstance(block, TypedTableBlock)]
        if not table_blocks:
            raise ValueError("structured-table recipe requires at least one table block")
        sheets: list[tuple[str, list[list[str]]]] = []
        used_names: set[str] = set()
        for block in table_blocks:
            sheet_name = _safe_sheet_name(block.unit_id, used_names)
            headers = ["row_key", *block.columns]
            values = [headers]
            seen: set[str] = set()
            expected = list(block.expected_row_keys)
            expected_index = {key: index for index, key in enumerate(expected)}
            rows = list(block.rows)
            if expected:
                rows.sort(key=lambda row: (expected_index.get(row.row_key, len(expected)), row.row_key))
            for row in rows:
                if not row.row_key:
                    raise ValueError("structured XLSX rows require non-empty row keys")
                if row.row_key in seen:
                    raise ValueError("duplicate table row keys are not allowed")
                if expected and row.row_key not in expected_index:
                    raise ValueError("table row key is outside the frozen expected scope")
                seen.add(row.row_key)
                row_values = [row.row_key, *[row.cells.get(column, "") for column in block.columns]]
                if any(_formula_like(value) for value in row_values):
                    raise ValueError("formula-like text is not allowed in generated table content")
                values.append([sanitize_xml10_text(value) for value in row_values])
            if expected and seen != set(expected):
                raise ValueError("structured XLSX table is missing an expected row")
            sheets.append((sheet_name, values))

        rendered = _xlsx_package(sheets)
        validate_ooxml_package(rendered, "xlsx")
        with zipfile.ZipFile(io.BytesIO(rendered), "r") as package:
            names = package.namelist()
            if any("vba" in name.lower() or "external" in name.lower() for name in names):
                raise ValueError("template-free XLSX must not contain active content")
        return self._result(
            rendered,
            model=document_model,
            recipe=recipe,
            binding=binding,
            profile=profile,
            preview={"sheet_count": len(sheets), "table_count": len(table_blocks)},
        )


def _heading_level(block: SectionBlock) -> int:
    value = (block.layout_hints or {}).get("level", 1)
    try:
        return max(1, min(9, int(value)))
    except (TypeError, ValueError):
        return 1


def _pdf_text(value: Any) -> str:
    normalized = sanitize_xml10_text(value).replace("\r\n", "\n").replace("\r", "\n")
    return escape(normalized).replace("\n", "<br/>")


def _formula_like(value: Any) -> bool:
    return str(value or "").startswith(_FORMULA_PREFIXES)


def _block_size(block: Any) -> tuple[int, int, int]:
    if getattr(block, "kind", "") == "table":
        return (
            sum(len(str(value)) for row in block.rows for value in row.cells.values()),
            len(block.rows),
            len(block.columns),
        )
    if getattr(block, "kind", "") == "list":
        return sum(len(str(value)) for value in block.items), 0, 0
    if getattr(block, "kind", "") == "section":
        return len(str(block.title or "")) + len(str(block.content or "")), 0, 0
    if getattr(block, "kind", "") == "paragraph":
        return len(str(block.content)), 0, 0
    if getattr(block, "kind", "") == "cross_reference":
        return sum(len(str(value)) for value in block.references), 0, 0
    return 0, 0, 0


def _canonical_zip(content: bytes) -> bytes:
    """Normalize ZIP metadata so identical semantic documents hash identically."""

    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9,
    ) as target:
        for name in sorted(source.namelist()):
            info = source.getinfo(name)
            normalized = zipfile.ZipInfo(filename=name, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.create_system = 0
            normalized.external_attr = 0
            normalized.comment = b""
            target.writestr(normalized, source.read(info))
    return output.getvalue()


def _safe_sheet_name(value: str, used: set[str]) -> str:
    import re

    base = re.sub(r"[\\/*?:\[\]]", "-", sanitize_xml10_text(value or "Report")).strip() or "Report"
    base = base[:31]
    result = base
    suffix = 1
    while result in used:
        suffix_text = f"-{suffix}"
        result = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used.add(result)
    return result


def _column_name(index: int) -> str:
    result = ""
    value = index + 1
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _xlsx_cell(reference: str, value: Any) -> str:
    text = escape(sanitize_xml10_text(value))
    preserve = ' xml:space="preserve"' if text[:1].isspace() or text[-1:].isspace() else ""
    return f'<c r="{reference}" t="inlineStr"><is><t{preserve}>{text}</t></is></c>'


def _xlsx_sheet(rows: list[list[str]]) -> str:
    rendered_rows: list[str] = []
    for row_index, row in enumerate(rows, start=1):
        cells = "".join(
            _xlsx_cell(f"{_column_name(column_index)}{row_index}", value)
            for column_index, value in enumerate(row)
        )
        rendered_rows.append(f'<row r="{row_index}">{cells}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<sheetData>" + "".join(rendered_rows) + "</sheetData></worksheet>"
    )


def _xlsx_package(sheets: list[tuple[str, list[list[str]]]]) -> bytes:
    content_types = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>',
    ]
    for index in range(1, len(sheets) + 1):
        content_types.append(
            f'<Override PartName="/xl/worksheets/sheet{index}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
    content_types.append("</Types>")
    workbook_sheets = []
    workbook_rels = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
    ]
    for index, (name, _rows) in enumerate(sheets, start=1):
        workbook_sheets.append(
            f'<sheet name="{escape(name)}" sheetId="{index}" '
            f'r:id="rId{index}"/>'
        )
        workbook_rels.append(
            f'<Relationship Id="rId{index}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            f'Target="worksheets/sheet{index}.xml"/>'
        )
    workbook_rels.append("</Relationships>")
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        "<sheets>" + "".join(workbook_sheets) + "</sheets></workbook>"
    )
    package_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as archive:
        archive.writestr("[Content_Types].xml", "".join(content_types))
        archive.writestr("_rels/.rels", package_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", "".join(workbook_rels))
        for index, (_name, rows) in enumerate(sheets, start=1):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", _xlsx_sheet(rows))
    return _canonical_zip(raw.getvalue())


__all__ = [
    "StructuredDocxRenderer",
    "StructuredPdfRenderer",
    "StructuredRenderResult",
    "StructuredXlsxRenderer",
]
