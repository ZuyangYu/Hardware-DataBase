"""Governed Artifact revision requests.

An Artifact is immutable.  A revision therefore records the parent artifact,
the exact input snapshot and the impacted validation scope before any future
writer/worker is allowed to create a child artifact.  This module does not
invent document bytes from a free-form request.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
import uuid
from collections.abc import Callable
from contextlib import closing
from datetime import datetime, timezone
from typing import Any, Mapping

from pydantic import BaseModel, Field, model_validator

from src.document_authoring.models import content_hash


REVISION_REVALIDATION_STATUSES = frozenset({"pending", "passed", "failed", "requires_human"})
_REVISION_STATUS_BY_REVALIDATION = {
    "passed": "revalidated",
    "failed": "failed",
    "requires_human": "waiting_human",
}


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _json(value: Any) -> str:
    if isinstance(value, BaseModel):
        value = value.model_dump(mode="json")
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _normalize_list(values: Any) -> list[str]:
    if values is None:
        return []
    if isinstance(values, str):
        values = [values]
    if not isinstance(values, (list, tuple, set)):
        raise ValueError("revision impact values must be a list")
    return list(dict.fromkeys(str(item).strip() for item in values if str(item).strip()))


_DIFF_CHANGE_LIMIT = 200
_DIFF_TEXT_LIMIT = 500


def _workbook_value_map(content: bytes) -> dict[str, str]:
    """Flat ``Sheet!Cell`` -> text map from workbook bytes (stdlib OOXML read)."""
    import zipfile
    from xml.etree import ElementTree as ET

    ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    ns_office_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    ns_pkg_rel = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    values: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        names = set(package.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(package.read("xl/sharedStrings.xml"))
            for entry in root.findall(f"{ns_main}si"):
                shared.append("".join(node.text or "" for node in entry.iter(f"{ns_main}t")))
        rels: dict[str, str] = {}
        if "xl/_rels/workbook.xml.rels" in names:
            rel_root = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
            for rel in rel_root.findall(f"{ns_pkg_rel}Relationship"):
                target = rel.get("Target") or ""
                if target.startswith("/"):
                    target = target[1:]
                elif target and not target.startswith("xl/"):
                    target = f"xl/{target}"
                rels[rel.get("Id") or ""] = target
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        for sheet in workbook.iter(f"{ns_main}sheet"):
            name = sheet.get("name") or ""
            target = rels.get(sheet.get(f"{ns_office_rel}id") or "", "")
            if not name or not target or target not in names:
                continue
            sheet_root = ET.fromstring(package.read(target))
            for cell in sheet_root.iter(f"{ns_main}c"):
                reference = cell.get("r") or ""
                if not reference:
                    continue
                cell_type = cell.get("t")
                value_node = cell.find(f"{ns_main}v")
                inline_node = cell.find(f"{ns_main}is")
                if cell_type == "inlineStr" and inline_node is not None:
                    text = "".join(node.text or "" for node in inline_node.iter(f"{ns_main}t"))
                elif cell_type == "s" and value_node is not None:
                    try:
                        text = shared[int(value_node.text or "0")]
                    except (ValueError, IndexError):
                        text = value_node.text or ""
                elif value_node is not None:
                    text = value_node.text or ""
                else:
                    text = ""
                values[f"{name}!{reference}"] = text
    return values


def _docx_value_map(content: bytes) -> dict[str, str]:
    """Ordered block -> text map from DOCX bytes (paragraphs and table cells)."""
    from docx import Document

    document = Document(io.BytesIO(content))
    values: dict[str, str] = {}
    for index, paragraph in enumerate(document.paragraphs):
        values[f"paragraph:{index:04d}"] = paragraph.text or ""
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                values[f"table:{table_index:02d}!r{row_index:03d}c{column_index:03d}"] = cell.text or ""
    return values


def _artifact_value_map(content: bytes, fmt: str) -> dict[str, str]:
    if fmt in {"xlsx", "xlsm"}:
        return _workbook_value_map(content)
    if fmt == "docx":
        return _docx_value_map(content)
    raise ValueError(f"unsupported artifact diff format: {fmt or 'unknown'}")


class ArtifactRevision(BaseModel):
    revision_id: str = Field(default_factory=lambda: f"artifact-revision-{uuid.uuid4().hex}")
    task_id: str
    work_order_id: str
    parent_artifact_id: str
    child_artifact_id: str | None = None
    status: str = "planned"
    request_type: str
    request: str
    changed_fields: list[str] = Field(default_factory=list)
    changed_sections: list[str] = Field(default_factory=list)
    input_snapshot_hash: str
    source_snapshot_hash: str
    template_version_id: str = ""
    schema_hash: str = ""
    parent_plan_id: str | None = None
    parent_plan_version: int | None = None
    parent_plan_hash: str | None = None
    child_plan_id: str | None = None
    child_plan_version: int | None = None
    child_plan_hash: str | None = None
    plan_diff_hash: str | None = None
    strategy_hash: str | None = None
    policy_hash: str | None = None
    affected_unit_ids: list[str] = Field(default_factory=list)
    reused_unit_ids: list[str] = Field(default_factory=list)
    execution_scope_hash: str | None = None
    impact_scope: dict[str, Any] = Field(default_factory=dict)
    revalidation_scope: list[str] = Field(default_factory=list)
    revalidation_status: str = "pending"
    revalidation_result: dict[str, Any] = Field(default_factory=dict)
    invalidated_approval_event_ids: list[str] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)
    client_request_id: str | None = None
    created_at: datetime = Field(default_factory=_utc_now)
    updated_at: datetime = Field(default_factory=_utc_now)

    @model_validator(mode="after")
    def validate_identity(self) -> ArtifactRevision:
        for field_name in (
            "revision_id", "task_id", "work_order_id", "parent_artifact_id",
            "request_type", "request", "input_snapshot_hash", "source_snapshot_hash",
            "status", "revalidation_status",
        ):
            value = str(getattr(self, field_name) or "").strip()
            if not value:
                raise ValueError(f"revision {field_name} is required")
            setattr(self, field_name, value)
        if self.revalidation_status not in REVISION_REVALIDATION_STATUSES:
            raise ValueError("unsupported revision revalidation status")
        self.changed_fields = _normalize_list(self.changed_fields)
        self.changed_sections = _normalize_list(self.changed_sections)
        self.affected_unit_ids = _normalize_list(self.affected_unit_ids)
        self.reused_unit_ids = _normalize_list(self.reused_unit_ids)
        self.revalidation_scope = _normalize_list(self.revalidation_scope)
        self.invalidated_approval_event_ids = _normalize_list(self.invalidated_approval_event_ids)
        plan_fields = (
            "parent_plan_id", "parent_plan_version", "parent_plan_hash",
        )
        if any(getattr(self, field_name) is not None for field_name in plan_fields) and not all(
            getattr(self, field_name) not in (None, "") for field_name in plan_fields
        ):
            raise ValueError("revision parent plan binding must include id, version and hash")
        child_fields = ("child_plan_id", "child_plan_version", "child_plan_hash")
        if any(getattr(self, field_name) is not None for field_name in child_fields) and not all(
            getattr(self, field_name) not in (None, "") for field_name in child_fields
        ):
            raise ValueError("revision child plan binding must include id, version and hash")
        if self.client_request_id is not None:
            self.client_request_id = str(self.client_request_id).strip() or None
        return self

    def to_dict(self) -> dict[str, Any]:
        return self.model_dump(mode="json")


Revision = ArtifactRevision


class ArtifactRevisionStore:
    """SQLite repository with request-level idempotency."""

    def __init__(self, db_path: str | os.PathLike[str]):
        self.db_path = os.fspath(db_path)
        if self.db_path != ":memory:":
            os.makedirs(os.path.dirname(self.db_path) or ".", exist_ok=True)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path, timeout=30, isolation_level=None)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        return conn

    def _init_db(self) -> None:
        with closing(self._connect()) as conn:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS artifact_revisions (
                    revision_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    work_order_id TEXT NOT NULL,
                    parent_artifact_id TEXT NOT NULL,
                    child_artifact_id TEXT,
                    status TEXT NOT NULL,
                    client_request_id TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    payload_json TEXT NOT NULL
                );
                CREATE UNIQUE INDEX IF NOT EXISTS idx_artifact_revisions_request
                    ON artifact_revisions(task_id, client_request_id)
                    WHERE client_request_id IS NOT NULL AND client_request_id != '';
                CREATE INDEX IF NOT EXISTS idx_artifact_revisions_task
                    ON artifact_revisions(task_id, created_at, revision_id);
                CREATE INDEX IF NOT EXISTS idx_artifact_revisions_parent
                    ON artifact_revisions(parent_artifact_id, created_at);
                """
            )

    @staticmethod
    def _from_row(row: sqlite3.Row | None) -> ArtifactRevision | None:
        if row is None:
            return None
        return ArtifactRevision.model_validate(json.loads(row["payload_json"]))

    @staticmethod
    def _identity(revision: ArtifactRevision) -> dict[str, Any]:
        payload = revision.to_dict()
        return {
            key: payload[key]
            for key in (
                "task_id", "work_order_id", "parent_artifact_id", "request_type", "request",
                "changed_fields", "changed_sections", "input_snapshot_hash", "source_snapshot_hash",
                "template_version_id", "schema_hash", "impact_scope", "revalidation_scope",
                "metadata", "client_request_id", "parent_plan_id", "parent_plan_version",
                "parent_plan_hash", "child_plan_id", "child_plan_version", "child_plan_hash",
                "plan_diff_hash", "strategy_hash", "policy_hash", "affected_unit_ids",
                "reused_unit_ids", "execution_scope_hash",
            )
        }

    def create(
        self,
        revision: ArtifactRevision | Mapping[str, Any] | None = None,
        **fields: Any,
    ) -> ArtifactRevision:
        if revision is not None and fields:
            raise TypeError("provide either a revision object or revision fields, not both")
        current = (
            revision
            if isinstance(revision, ArtifactRevision)
            else ArtifactRevision.model_validate(revision)
            if revision is not None
            else ArtifactRevision(**fields)
        )
        identity = self._identity(current)
        with closing(self._connect()) as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                existing_row = None
                if current.client_request_id:
                    existing_row = conn.execute(
                        """SELECT * FROM artifact_revisions
                           WHERE task_id = ? AND client_request_id = ?""",
                        (current.task_id, current.client_request_id),
                    ).fetchone()
                if existing_row is not None:
                    existing = self._from_row(existing_row)
                    assert existing is not None
                    if self._identity(existing) != identity:
                        raise ValueError("revision idempotency key conflicts with existing payload")
                    conn.execute("COMMIT")
                    return existing
                collision = conn.execute(
                    "SELECT 1 FROM artifact_revisions WHERE revision_id = ?",
                    (current.revision_id,),
                ).fetchone()
                if collision is not None:
                    raise ValueError("revision id already exists")
                conn.execute(
                    """INSERT INTO artifact_revisions (
                           revision_id, task_id, work_order_id, parent_artifact_id,
                           child_artifact_id, status, client_request_id, created_at,
                           updated_at, payload_json
                       ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        current.revision_id, current.task_id, current.work_order_id,
                        current.parent_artifact_id, current.child_artifact_id, current.status,
                        current.client_request_id, current.created_at.isoformat(),
                        current.updated_at.isoformat(), _json(current),
                    ),
                )
                conn.execute("COMMIT")
                return current
            except Exception:
                conn.execute("ROLLBACK")
                raise

    def get(self, revision_id: str) -> ArtifactRevision | None:
        with closing(self._connect()) as conn:
            row = conn.execute(
                "SELECT * FROM artifact_revisions WHERE revision_id = ?",
                (str(revision_id or "").strip(),),
            ).fetchone()
        return self._from_row(row)

    def list_for_task(self, task_id: str, *, limit: int = 100) -> list[ArtifactRevision]:
        with closing(self._connect()) as conn:
            rows = conn.execute(
                """SELECT * FROM artifact_revisions
                   WHERE task_id = ? ORDER BY created_at, revision_id LIMIT ?""",
                (str(task_id or "").strip(), max(1, min(int(limit), 200))),
            ).fetchall()
        return [item for row in rows if (item := self._from_row(row)) is not None]

    def update_status(self, revision_id: str, status: str) -> ArtifactRevision:
        normalized = str(status or "").strip()
        if not normalized:
            raise ValueError("revision status is required")
        current = self.get(revision_id)
        if current is None:
            raise KeyError("artifact revision not found")
        revised = current.model_copy(update={"status": normalized, "updated_at": _utc_now()})
        with closing(self._connect()) as conn:
            conn.execute(
                "UPDATE artifact_revisions SET status = ?, updated_at = ?, payload_json = ? WHERE revision_id = ?",
                (revised.status, revised.updated_at.isoformat(), _json(revised), revised.revision_id),
            )
        return revised

    def bind_child_artifact(self, revision_id: str, child_artifact_id: str) -> ArtifactRevision:
        child_id = str(child_artifact_id or "").strip()
        if not child_id:
            raise ValueError("child artifact id is required")
        current = self.get(revision_id)
        if current is None:
            raise KeyError("artifact revision not found")
        if current.child_artifact_id and current.child_artifact_id != child_id:
            raise ValueError("artifact revision is already bound to another child artifact")
        revised = current.model_copy(update={
            "child_artifact_id": child_id,
            "status": "completed",
            "updated_at": _utc_now(),
        })
        with closing(self._connect()) as conn:
            conn.execute(
                "UPDATE artifact_revisions SET child_artifact_id = ?, status = ?, updated_at = ?, payload_json = ? WHERE revision_id = ?",
                (child_id, revised.status, revised.updated_at.isoformat(), _json(revised), revised.revision_id),
            )
        return revised

    def record_revalidation(
        self,
        revision_id: str,
        child_artifact_id: str,
        *,
        revalidation_status: str,
        revalidation_result: Mapping[str, Any] | None = None,
    ) -> ArtifactRevision:
        """Bind one generated child and its immutable revalidation result.

        The writer/worker owns creating the child bytes.  This repository only
        commits the result of that controlled operation and makes an identical
        completion callback safe to replay.
        """
        child_id = str(child_artifact_id or "").strip()
        normalized_status = str(revalidation_status or "").strip().casefold()
        if not child_id:
            raise ValueError("child artifact id is required")
        if normalized_status not in _REVISION_STATUS_BY_REVALIDATION:
            raise ValueError("revalidation status must be passed, failed or requires_human")
        if revalidation_result is not None and not isinstance(revalidation_result, Mapping):
            raise ValueError("revalidation result must be an object")
        result = dict(revalidation_result or {})
        with closing(self._connect()) as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                row = conn.execute(
                    "SELECT * FROM artifact_revisions WHERE revision_id = ?",
                    (str(revision_id or "").strip(),),
                ).fetchone()
                current = self._from_row(row)
                if current is None:
                    raise KeyError("artifact revision not found")
                if current.child_artifact_id and current.child_artifact_id != child_id:
                    raise ValueError("artifact revision is already bound to another child artifact")
                if current.revalidation_status != "pending":
                    if (
                        current.revalidation_status != normalized_status
                        or current.revalidation_result != result
                    ):
                        raise ValueError("revalidation result conflicts with existing revision completion")
                    conn.execute("COMMIT")
                    return current
                revised = current.model_copy(update={
                    "child_artifact_id": child_id,
                    "status": _REVISION_STATUS_BY_REVALIDATION[normalized_status],
                    "revalidation_status": normalized_status,
                    "revalidation_result": result,
                    "updated_at": _utc_now(),
                })
                conn.execute(
                    """UPDATE artifact_revisions SET
                           child_artifact_id = ?, status = ?, updated_at = ?, payload_json = ?
                       WHERE revision_id = ?""",
                    (
                        child_id, revised.status, revised.updated_at.isoformat(),
                        _json(revised), revised.revision_id,
                    ),
                )
                conn.execute("COMMIT")
                return revised
            except Exception:
                conn.execute("ROLLBACK")
                raise


RevisionStore = ArtifactRevisionStore


class DocumentRevisionService:
    """Authorize and persist a revision request without fabricating content."""

    def __init__(
        self,
        *,
        authoring_store: Any,
        task_store: Any,
        revision_store: ArtifactRevisionStore,
        source_snapshot_resolver: Callable[[Any], Any] | None = None,
    ):
        self.authoring_store = authoring_store
        self.task_store = task_store
        self.store = revision_store
        self.source_snapshot_resolver = source_snapshot_resolver

    def create_revision(
        self,
        ctx: Any,
        *,
        task_id: str,
        parent_artifact_id: str,
        request_type: str,
        request: str,
        changed_fields: Any = None,
        changed_sections: Any = None,
        client_request_id: str,
        metadata: Mapping[str, Any] | None = None,
        parent_plan: Any | None = None,
        child_plan: Any | None = None,
        plan_diff: Any | None = None,
        strategy_hash: str | None = None,
        policy_hash: str | None = None,
    ) -> ArtifactRevision:
        task = self._task_for_context(ctx, task_id, required_permission="write")
        parent_id = str(parent_artifact_id or "").strip()
        if not parent_id:
            raise ValueError("parent artifact id is required")
        parent = self.authoring_store.get_artifact(parent_id)
        if parent is None:
            raise KeyError("parent artifact not found")
        work_order_id = str(getattr(parent, "work_order_id", "") or "").strip()
        order = self.authoring_store.get_work_order(work_order_id) if work_order_id else None
        if order is None:
            raise KeyError("parent artifact work order not found")
        order_tenant = str(getattr(order, "tenant_id", "") or "").strip()
        if order_tenant and order_tenant != task.tenant_id:
            raise PermissionError("parent artifact work order is outside the current tenant")
        order_kb = str(getattr(order, "knowledge_base_name", "") or "").strip()
        if order_kb and order_kb != str(task.knowledge_base_name or "").strip():
            raise PermissionError("parent artifact work order is outside the current knowledge base")
        if getattr(order, "task_id", None) not in {None, task.task_id}:
            raise ValueError("parent artifact work order does not belong to task")
        if task.work_order_id and task.work_order_id != work_order_id:
            raise ValueError("parent artifact is not part of the task's current work order")

        request_kind = str(request_type or "").strip().lower()
        if request_kind not in {"field_update", "section_update", "full_regeneration"}:
            raise ValueError("unsupported artifact revision request type")
        normalized_request = str(request or "").strip()
        if not normalized_request:
            raise ValueError("revision request is required")
        fields = _normalize_list(changed_fields)
        sections = _normalize_list(changed_sections)
        if request_kind == "field_update" and not fields:
            raise ValueError("field_update requires changed_fields")
        if request_kind == "section_update" and not sections:
            raise ValueError("section_update requires changed_sections")
        if request_kind == "full_regeneration":
            fields = []
            sections = []

        source_snapshot_hash = ""
        if self.source_snapshot_resolver is not None:
            snapshot = self.source_snapshot_resolver(order)
            source_snapshot_hash = str(getattr(snapshot, "content_hash", "") or "").strip()
            if not source_snapshot_hash:
                raise ValueError("resolved source snapshot has no content hash")
        if not source_snapshot_hash:
            source_snapshot_hash = (
                str(getattr(order, "baseline_content_hash", "") or "").strip()
                or str(getattr(order, "source_set_snapshot_id", "") or "").strip()
            )
        if not source_snapshot_hash:
            source_snapshot_hash = content_hash({"work_order_id": work_order_id})
        schema_hash = content_hash({
            "template_version_id": getattr(order, "template_version_id", None),
            "template_schema_id": getattr(order, "template_schema_id", None),
            "template_schema_version": getattr(order, "template_schema_version", None),
            "document_schema_id": getattr(order, "document_schema_id", None),
            "document_schema_version": getattr(order, "document_schema_version", None),
        })
        input_snapshot_hash = content_hash({
            "parent_artifact_id": parent_id,
            "parent_artifact_content_hash": getattr(parent, "content_hash", None),
            "validation_report_id": getattr(parent, "validation_report_id", None),
            "source_snapshot_hash": source_snapshot_hash,
            "schema_hash": schema_hash,
        })
        parent_plan_id = parent_plan_version = parent_plan_hash = None
        child_plan_id = child_plan_version = child_plan_hash = None
        plan_diff_hash = None
        affected_unit_ids: list[str] = []
        reused_unit_ids: list[str] = []
        execution_scope_hash = None
        frozen_strategy_hash = str(strategy_hash or "").strip() or None
        frozen_policy_hash = str(policy_hash or "").strip() or None
        if parent_plan is not None or child_plan is not None or plan_diff is not None:
            from src.document_authoring.planning.diff import compile_affected_subgraph, diff_document_plans
            from src.document_authoring.planning.models import DocumentPlan, PlanDiff

            if parent_plan is None or child_plan is None:
                raise ValueError("plan-bound revisions require both parent_plan and child_plan")
            frozen_parent_plan = (
                parent_plan if isinstance(parent_plan, DocumentPlan)
                else DocumentPlan.model_validate(parent_plan)
            )
            frozen_child_plan = (
                child_plan if isinstance(child_plan, DocumentPlan)
                else DocumentPlan.model_validate(child_plan)
            )
            frozen_diff = (
                plan_diff if isinstance(plan_diff, PlanDiff)
                else PlanDiff.model_validate(plan_diff)
                if plan_diff is not None
                else diff_document_plans(frozen_parent_plan, frozen_child_plan)
            )
            if (
                frozen_diff.parent_plan_id != frozen_parent_plan.document_plan_id
                or frozen_diff.parent_plan_version != frozen_parent_plan.version
                or frozen_diff.child_plan_id != frozen_child_plan.document_plan_id
                or frozen_diff.child_plan_version != frozen_child_plan.version
            ):
                raise ValueError("plan diff is not bound to the supplied parent and child plans")
            if frozen_parent_plan.source_snapshot_hash != frozen_child_plan.source_snapshot_hash:
                raise ValueError("plan-bound revision cannot mix source snapshots")
            if source_snapshot_hash != frozen_parent_plan.source_snapshot_hash:
                raise ValueError("revision source snapshot does not match the parent plan")
            scope = compile_affected_subgraph(frozen_parent_plan, frozen_child_plan, frozen_diff)
            parent_plan_id = frozen_parent_plan.document_plan_id
            parent_plan_version = frozen_parent_plan.version
            parent_plan_hash = frozen_parent_plan.plan_hash
            child_plan_id = frozen_child_plan.document_plan_id
            child_plan_version = frozen_child_plan.version
            child_plan_hash = frozen_child_plan.plan_hash
            plan_diff_hash = frozen_diff.diff_hash
            affected_unit_ids = list(scope.affected_unit_ids)
            reused_unit_ids = list(scope.reused_unit_ids)
            execution_scope_hash = scope.scope_hash
            frozen_strategy_hash = frozen_strategy_hash or content_hash({
                "strategy_id": frozen_child_plan.domain_strategy_id,
                "strategy_version": frozen_child_plan.domain_strategy_version,
            })
            frozen_policy_hash = frozen_policy_hash or content_hash({
                "unit_review_policy": frozen_child_plan.unit_review_policy,
                "document_review_policy": frozen_child_plan.document_review_policy,
                "approval_policy": frozen_child_plan.approval_policy,
            })
        impact_kind = (
            "full"
            if request_kind == "full_regeneration" or not fields and not sections
            else "mixed"
            if fields and sections
            else "fields"
            if fields
            else "sections"
        )
        impact_scope = {
            "kind": impact_kind,
            "fields": fields,
            "sections": sections,
            "requires_full_revalidation": impact_kind == "full",
        }
        if parent_plan_id is not None:
            revalidation_scope = [
                "artifact", "approval", "document_review", "render", *fields, *sections,
                *affected_unit_ids,
            ]
        else:
            # Preserve the historical revision projection for legacy rows.
            revalidation_scope = ["artifact", "approval", *fields, *sections]
        revision = self.store.create(
            task_id=task.task_id,
            work_order_id=work_order_id,
            parent_artifact_id=parent_id,
            status="planned",
            request_type=request_kind,
            request=normalized_request,
            changed_fields=fields,
            changed_sections=sections,
            input_snapshot_hash=input_snapshot_hash,
            source_snapshot_hash=source_snapshot_hash,
            template_version_id=str(getattr(order, "template_version_id", "") or ""),
            schema_hash=schema_hash,
            parent_plan_id=parent_plan_id,
            parent_plan_version=parent_plan_version,
            parent_plan_hash=parent_plan_hash,
            child_plan_id=child_plan_id,
            child_plan_version=child_plan_version,
            child_plan_hash=child_plan_hash,
            plan_diff_hash=plan_diff_hash,
            strategy_hash=frozen_strategy_hash,
            policy_hash=frozen_policy_hash,
            affected_unit_ids=affected_unit_ids,
            reused_unit_ids=reused_unit_ids,
            execution_scope_hash=execution_scope_hash,
            impact_scope=impact_scope,
            revalidation_scope=revalidation_scope,
            invalidated_approval_event_ids=list(getattr(parent, "approval_event_ids", []) or []),
            metadata=dict(metadata or {}),
            client_request_id=str(client_request_id or "").strip(),
        )
        update_artifact = getattr(self.authoring_store, "update_artifact", None)
        if callable(update_artifact):
            update_artifact(
                parent_id,
                validity_status="revalidation_required",
                regeneration_status="recommended",
            )
        update_task_status = getattr(self.task_store, "update_status", None)
        if callable(update_task_status):
            update_task_status(task.task_id, "waiting_human")
        return revision

    def get_revision(self, ctx: Any, revision_id: str) -> ArtifactRevision | None:
        revision = self.store.get(revision_id)
        if revision is None:
            return None
        self._task_for_context(ctx, revision.task_id, required_permission="read")
        return revision

    def complete_revision(
        self,
        ctx: Any,
        revision_id: str,
        *,
        child_artifact_id: str,
        revalidation_status: str,
        revalidation_result: Mapping[str, Any] | None = None,
    ) -> ArtifactRevision:
        """Finalize a governed revision after a worker produced its child.

        This method intentionally accepts an artifact reference rather than
        document bytes.  Generation remains owned by the existing controlled
        writer/worker; the completion boundary verifies lineage, records the
        validation outcome and keeps the parent immutable.
        """
        revision = self.store.get(revision_id)
        if revision is None:
            raise KeyError("artifact revision not found")
        task = self._task_for_context(ctx, revision.task_id, required_permission="write")
        if revision.task_id != task.task_id:
            raise PermissionError("artifact revision is outside the current task")
        normalized_status = str(revalidation_status or "").strip().casefold()
        if normalized_status not in _REVISION_STATUS_BY_REVALIDATION:
            raise ValueError("revalidation status must be passed, failed or requires_human")
        result = dict(revalidation_result or {})
        child_id = str(child_artifact_id or "").strip()
        if not child_id:
            raise ValueError("child artifact id is required")
        child = self.authoring_store.get_artifact(child_id)
        if child is None:
            raise KeyError("child artifact not found")
        if child.artifact_id == revision.parent_artifact_id:
            raise ValueError("revision child artifact must differ from its parent")
        if not self._child_in_revision_lineage(revision, child):
            raise ValueError("revision child artifact does not belong to the revision work order")
        child_tenant = str(getattr(child, "tenant_id", "") or "").strip()
        if child_tenant and child_tenant != task.tenant_id:
            raise PermissionError("revision child artifact is outside the current tenant")
        if str(getattr(child, "parent_artifact_id", "") or "") != revision.parent_artifact_id:
            raise ValueError("revision child artifact is not derived from the revision parent")
        if str(getattr(child, "stage", "") or "") == "approved_release":
            raise ValueError("revision child must be reviewed before release")
        for key, expected in {
            "input_snapshot_hash": revision.input_snapshot_hash,
            "source_snapshot_hash": revision.source_snapshot_hash,
            "schema_hash": revision.schema_hash,
            "child_artifact_content_hash": str(getattr(child, "content_hash", "") or ""),
        }.items():
            supplied = result.get(key)
            if supplied is not None and str(supplied) != str(expected):
                raise ValueError(f"revalidation result {key} does not match the revision snapshot")
            result[key] = expected
        diff_report = self._diff_report_for(revision, child)
        if diff_report is not None:
            result["diff_report"] = diff_report
        return self._apply_child_binding(revision, child, normalized_status, result)

    def _child_in_revision_lineage(self, revision: ArtifactRevision, child: Any) -> bool:
        """True when the child was produced by the revision work order itself
        or by a restart descendant of it (revision regeneration chain)."""
        child_order_id = str(getattr(child, "work_order_id", "") or "")
        if child_order_id == revision.work_order_id:
            return True
        if not hasattr(self.authoring_store, "get_work_order"):
            return False
        current_id = child_order_id
        for _hop in range(32):
            if not current_id:
                return False
            order = self.authoring_store.get_work_order(current_id)
            if order is None:
                return False
            parent_order_id = str(getattr(order, "restart_of_work_order_id", "") or "")
            if not parent_order_id:
                return False
            if parent_order_id == revision.work_order_id:
                return True
            current_id = parent_order_id
        return False

    def _diff_report_for(self, revision: ArtifactRevision, child: Any) -> dict[str, Any] | None:
        """Structured parent -> child diff persisted inside the revision.

        The report is deterministic (sorted locations, bounded change list) so
        idempotent replays compare equal.  A diff failure must never block the
        governed binding, so failures degrade to an ``error`` entry.
        """
        reader = getattr(self.authoring_store, "read_artifact_content", None)
        if not callable(reader):
            return None
        parent_id = revision.parent_artifact_id
        child_id = str(getattr(child, "artifact_id", "") or "")
        try:
            parent = self.authoring_store.get_artifact(parent_id)
            parent_format = str(getattr(parent, "output_format", "") or "") if parent is not None else ""
            fmt = str(getattr(child, "output_format", "") or "") or parent_format
            before_map = _artifact_value_map(reader(parent_id), fmt)
            after_map = _artifact_value_map(reader(child_id), fmt)
        except Exception as exc:  # pragma: no cover - storage specific failures
            return {
                "parent_artifact_id": parent_id,
                "child_artifact_id": child_id,
                "error": f"diff unavailable: {exc}",
            }
        changes: list[dict[str, Any]] = []
        changed = added = removed = 0
        unchanged = 0
        for location in sorted(set(before_map) | set(after_map)):
            before = before_map.get(location, "")
            after = after_map.get(location, "")
            if before == after:
                unchanged += 1
                continue
            if location not in before_map:
                kind = "added"
                added += 1
            elif location not in after_map:
                kind = "removed"
                removed += 1
            else:
                kind = "changed"
                changed += 1
            if len(changes) < _DIFF_CHANGE_LIMIT:
                changes.append({
                    "location": location,
                    "kind": kind,
                    "before": before[:_DIFF_TEXT_LIMIT],
                    "after": after[:_DIFF_TEXT_LIMIT],
                })
        return {
            "parent_artifact_id": parent_id,
            "child_artifact_id": child_id,
            "format": fmt or None,
            "summary": {
                "changed": changed, "added": added, "removed": removed, "unchanged": unchanged,
            },
            "changes": changes,
            "truncated": (changed + added + removed) > len(changes),
        }

    def bind_generated_child(
        self,
        order: Any,
        child: Any,
        *,
        revalidation_status: str,
        revalidation_result: Mapping[str, Any] | None = None,
    ) -> ArtifactRevision:
        """Complete a revision from the worker path, without a user context.

        The durable authoring worker runs without request identity.  It is the
        trusted executor for the work order it just finalized, so lineage
        (child belongs to the revision work order and descends from the
        revision parent) is checked directly against stored facts.
        """
        revision_id = str(getattr(order, "revision_id", None) or "").strip()
        revision = self.store.get(revision_id)
        if revision is None:
            raise KeyError("artifact revision not found")
        normalized_status = str(revalidation_status or "").strip().casefold()
        if normalized_status not in REVISION_REVALIDATION_STATUSES or normalized_status == "pending":
            raise ValueError("revalidation status must be passed, failed or requires_human")
        child_id = str(getattr(child, "artifact_id", "") or "").strip()
        if not child_id:
            raise ValueError("revision child artifact id is required")
        if child_id == revision.parent_artifact_id:
            raise ValueError("revision child artifact must differ from its parent")
        if not self._child_in_revision_lineage(revision, child):
            raise ValueError("revision child artifact does not belong to the revision work order")
        if str(getattr(child, "parent_artifact_id", "") or "") != revision.parent_artifact_id:
            raise ValueError("revision child artifact is not derived from the revision parent")
        if str(getattr(child, "stage", "") or "") == "approved_release":
            raise ValueError("revision child must be reviewed before release")
        result = dict(revalidation_result or {})
        result.update({
            "input_snapshot_hash": revision.input_snapshot_hash,
            "source_snapshot_hash": revision.source_snapshot_hash,
            "schema_hash": revision.schema_hash,
            "child_artifact_content_hash": str(getattr(child, "content_hash", "") or ""),
        })
        diff_report = self._diff_report_for(revision, child)
        if diff_report is not None:
            result["diff_report"] = diff_report
        return self._apply_child_binding(revision, child, normalized_status, result)

    def _apply_child_binding(
        self,
        revision: ArtifactRevision,
        child: Any,
        normalized_status: str,
        result: dict[str, Any],
    ) -> ArtifactRevision:
        child_id = str(getattr(child, "artifact_id", "") or "").strip()
        if revision.child_artifact_id:
            if revision.child_artifact_id != child_id:
                raise ValueError("artifact revision is already bound to another child artifact")
            if revision.revalidation_status != "pending":
                if (
                    revision.revalidation_status != normalized_status
                    or revision.revalidation_result != result
                ):
                    raise ValueError("revalidation result conflicts with existing revision completion")
                self._mark_task_waiting(revision.task_id)
                return revision

        update_artifact = getattr(self.authoring_store, "update_artifact", None)
        if callable(update_artifact):
            current_reasons = list(getattr(child, "status_reasons", []) or [])
            reason = {
                "code": "artifact_revision_revalidated",
                "revision_id": revision.revision_id,
                "revalidation_status": normalized_status,
                "revalidation_result_hash": content_hash(result),
            }
            if not any(
                isinstance(item, Mapping)
                and item.get("code") == reason["code"]
                and item.get("revision_id") == reason["revision_id"]
                and item.get("revalidation_status") == reason["revalidation_status"]
                for item in current_reasons
            ):
                current_reasons.append(reason)
            update_artifact(
                child_id,
                revision_id=revision.revision_id,
                validity_status="current" if normalized_status == "passed" else "revalidation_required",
                regeneration_status="not_needed" if normalized_status == "passed" else "recommended",
                status_reasons=current_reasons,
            )
        completed = self.store.record_revalidation(
            revision.revision_id,
            child_id,
            revalidation_status=normalized_status,
            revalidation_result=result,
        )
        self._mark_task_waiting(revision.task_id)
        return completed

    def _mark_task_waiting(self, task_id: str) -> None:
        update_task_status = getattr(self.task_store, "update_status", None)
        if callable(update_task_status):
            update_task_status(task_id, "waiting_human")

    def mark_released(self, ctx: Any, child_artifact_id: str) -> ArtifactRevision | None:
        """Close a revision after its child has passed the normal release gate."""
        child_id = str(child_artifact_id or "").strip()
        if not child_id:
            raise ValueError("child artifact id is required")
        child = self.authoring_store.get_artifact(child_id)
        if child is None:
            raise KeyError("child artifact not found")
        task_id = ""
        for revision in self.store.list_for_task(
            str(getattr(child, "task_id", "") or "")
        ) if getattr(child, "task_id", None) else []:
            if revision.child_artifact_id == child_id:
                task_id = revision.task_id
                break
        if not task_id:
            # The artifact store does not persist task_id on every historical
            # artifact.  Resolve the candidate through its work-order lineage.
            order = self.authoring_store.get_work_order(getattr(child, "work_order_id", ""))
            task_id = str(getattr(order, "task_id", "") or "")
        if not task_id:
            return None
        task = self._task_for_context(ctx, task_id, required_permission="write")
        candidates = [
            revision for revision in self.store.list_for_task(task.task_id)
            if revision.child_artifact_id == child_id
        ]
        if not candidates:
            return None
        revision = candidates[-1]
        if revision.status == "completed":
            return revision
        return self.store.update_status(revision.revision_id, "completed")

    def list_revisions(self, ctx: Any, task_id: str) -> list[ArtifactRevision]:
        task = self._task_for_context(ctx, task_id, required_permission="read")
        return self.store.list_for_task(task.task_id)

    def _task_for_context(self, ctx: Any, task_id: str, *, required_permission: str) -> Any:
        task = self.task_store.get(str(task_id or "").strip())
        if task is None:
            raise KeyError("document task not found")
        if task.tenant_id != str(getattr(ctx, "tenant_id", None) or "default"):
            raise PermissionError("document task is outside the current tenant")
        if task.user_id != str(getattr(ctx, "user_id", None) or ""):
            raise PermissionError("document task is outside the current owner scope")
        kb_name = str(task.knowledge_base_name or "").strip()
        if kb_name and not ctx.has_kb_permission(kb_name, required_permission):
            raise PermissionError(f"knowledge base {required_permission} permission is required")
        return task

__all__ = [
    "ArtifactRevision", "ArtifactRevisionStore", "DocumentRevisionService",
    "REVISION_REVALIDATION_STATUSES",
    "Revision", "RevisionStore",
]
